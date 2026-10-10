from __future__ import annotations

import io
import json
import os
import re
import shutil
import sqlite3
import uuid
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path
from typing import Any

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel

ROOT = Path(__file__).resolve().parent
PROJECT = ROOT.parent
DATA = Path(os.environ.get("WP_DATA_DIR", ROOT / "data")).resolve()
UPLOADS = DATA / "uploads"
DB_PATH = DATA / "wp_demo.db"
AMOS_DIR = PROJECT / "DATABASE"
SAMPLES_DIR = PROJECT / "TALLY SHEET KHACH HANG"
DATA.mkdir(exist_ok=True)
UPLOADS.mkdir(exist_ok=True)

app = FastAPI(title="VAECO Workpack Evaluation", version="1.0.0")
app.mount("/static", StaticFiles(directory=ROOT / "static"), name="static")


@contextmanager
def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys=ON")
    try:
        yield conn
        conn.commit()
    finally:
        conn.close()


def now() -> str:
    return datetime.now().isoformat(timespec="seconds")


def norm(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).upper()


def mhr_hours(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    text = str(value).strip()
    if ":" in text:
        try:
            hours, minutes = text.split(":", 1)
            return round(float(hours) + float(minutes) / 60, 2)
        except ValueError:
            return 0.0
    try:
        return round(float(text), 2)
    except ValueError:
        return 0.0


def canonical_aircraft(value: str) -> str:
    text = norm(value).replace("AIRBUS", "").replace("BOEING", "").replace("-", "").strip()
    aliases = {"A320": "320", "A321": "321", "A350": "350", "B787": "787", "ATR72": "ATR"}
    return aliases.get(text, text)


def init_db() -> None:
    with db() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS amp_task (
              id INTEGER PRIMARY KEY, source TEXT NOT NULL, aircraft_type TEXT,
              customer TEXT, task_no TEXT NOT NULL, task_key TEXT NOT NULL,
              title TEXT, description TEXT, phase TEXT, zone TEXT, area TEXT,
              panel TEXT, special_code TEXT, resource_type TEXT, classification TEXT,
              type_code TEXT, qualification TEXT, planning_mhr TEXT, information_skill TEXT,
              requirement_link TEXT, linked_jobcards TEXT, jobcards_title TEXT
            );
            CREATE INDEX IF NOT EXISTS idx_amp_key ON amp_task(task_key);
            CREATE INDEX IF NOT EXISTS idx_amp_customer ON amp_task(customer, task_key);
            CREATE TABLE IF NOT EXISTS jobcard_master (
              id INTEGER PRIMARY KEY, meta_key TEXT, meta_value TEXT,
              aircraft_type TEXT, jobcard_no TEXT NOT NULL, jobcard_key TEXT NOT NULL,
              title TEXT, description TEXT, phase TEXT, zone TEXT, area TEXT,
              panel TEXT, special_code TEXT, resource_type TEXT, classification TEXT,
              qualification TEXT, planning_mhr TEXT, information_skill TEXT,
              information_mhr TEXT, requirement_link TEXT
            );
            CREATE INDEX IF NOT EXISTS idx_jobcard_key ON jobcard_master(jobcard_key);
            CREATE TABLE IF NOT EXISTS external_data (
              id INTEGER PRIMARY KEY AUTOINCREMENT, category TEXT NOT NULL,
              code TEXT NOT NULL, canonical_value TEXT, aircraft_type TEXT,
              customer TEXT, description TEXT, reason TEXT,
              status TEXT DEFAULT 'Có hiệu lực', updated_by TEXT DEFAULT 'Quản trị dữ liệu',
              updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS tally_sheet (
              id TEXT PRIMARY KEY, file_name TEXT NOT NULL, stored_path TEXT NOT NULL,
              customer TEXT NOT NULL, aircraft_reg TEXT, aircraft_type TEXT,
              check_type TEXT, adapter TEXT NOT NULL, version INTEGER DEFAULT 1,
              status TEXT DEFAULT 'Đã xử lý', uploader TEXT DEFAULT 'Demo User',
              uploaded_at TEXT NOT NULL, row_count INTEGER DEFAULT 0,
              matched_count INTEGER DEFAULT 0, error_message TEXT
            );
            CREATE TABLE IF NOT EXISTS evaluation (
              id TEXT PRIMARY KEY, tally_id TEXT NOT NULL REFERENCES tally_sheet(id),
              status TEXT DEFAULT 'Đang review', revision TEXT, tat_review TEXT,
              assumptions TEXT, notes TEXT, prepared_by TEXT, approved_by TEXT,
              created_at TEXT NOT NULL, updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS evaluation_item (
              id INTEGER PRIMARY KEY AUTOINCREMENT, evaluation_id TEXT NOT NULL REFERENCES evaluation(id),
              seq TEXT, section TEXT, task_no TEXT, source_title TEXT, event_type TEXT,
              rii TEXT, source_note TEXT, matched INTEGER DEFAULT 0, amp_source TEXT,
              title TEXT, description TEXT, phase TEXT, zone TEXT, area TEXT, panel TEXT,
              special_code TEXT, resource_type TEXT, classification TEXT, qualification TEXT,
              planning_mhr REAL DEFAULT 0, reviewed_mhr REAL, information_skill TEXT,
              requirement_link TEXT, linked_jobcards TEXT, type_code TEXT, jobcards_title TEXT,
              important_job INTEGER DEFAULT 0,
              important_source TEXT, not_in_database INTEGER DEFAULT 0, review_note TEXT,
              review_reason TEXT, updated_at TEXT
            );
            CREATE INDEX IF NOT EXISTS idx_items_eval ON evaluation_item(evaluation_id);
            CREATE TABLE IF NOT EXISTS change_log (
              id INTEGER PRIMARY KEY AUTOINCREMENT, evaluation_id TEXT, item_id INTEGER,
              field_name TEXT, old_value TEXT, new_value TEXT, reason TEXT,
              changed_by TEXT, changed_at TEXT
            );
            CREATE TABLE IF NOT EXISTS workflow_step (
              tally_id TEXT NOT NULL REFERENCES tally_sheet(id) ON DELETE CASCADE,
              step_code TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'not_started',
              details_json TEXT NOT NULL DEFAULT '{}', completed_at TEXT,
              updated_at TEXT NOT NULL,
              PRIMARY KEY (tally_id, step_code)
            );
            CREATE UNIQUE INDEX IF NOT EXISTS idx_evaluation_tally ON evaluation(tally_id);
            """
        )
        # Keep existing demo databases compatible when the output specification gains columns.
        for table, column in (("amp_task", "type_code"), ("amp_task", "jobcards_title"),
                              ("evaluation_item", "type_code"), ("evaluation_item", "jobcards_title"),
                              ("tally_sheet", "description")):
            existing = {row[1] for row in conn.execute(f"PRAGMA table_info({table})")}
            if column not in existing:
                conn.execute(f"ALTER TABLE {table} ADD COLUMN {column} TEXT")
        if not conn.execute("SELECT COUNT(*) FROM external_data").fetchone()[0]:
            seed_rows = [
                ("Mapping khách hàng", "VA-VJ", "VietJet", None, "VA-VJ", "Mẫu Tally Sheet VietJet", "Chuẩn hóa mã khách hàng"),
                ("Mapping khách hàng", "DLH", "Lufthansa", None, "DLH", "Mẫu Workpackage Summary Lufthansa", "Chuẩn hóa mã khách hàng"),
                ("Mapping loại tàu bay", "A320", "320", "A320", None, "Airbus A320", "Chuẩn hóa định danh loại tàu bay"),
                ("Mapping loại tàu bay", "A321", "321", "A321", None, "Airbus A321", "Chuẩn hóa định danh loại tàu bay"),
                ("Mapping loại tàu bay", "A350", "350", "A350", None, "Airbus A350", "Chuẩn hóa định danh loại tàu bay"),
                ("Mapping loại tàu bay", "B787", "787", "B787", None, "Boeing 787", "Chuẩn hóa định danh loại tàu bay"),
            ]
            conn.executemany("""INSERT INTO external_data(category,code,canonical_value,aircraft_type,
              customer,description,reason,updated_at) VALUES (?,?,?,?,?,?,?,?)""",
              [(*row, now()) for row in seed_rows])


def cellmap(headers: list[Any]) -> dict[str, int]:
    return {norm(value): idx for idx, value in enumerate(headers) if value not in (None, "")}


def import_amos() -> None:
    with db() as conn:
        if conn.execute("SELECT COUNT(*) FROM amp_task").fetchone()[0]:
            return
    sources = [
        (AMOS_DIR / "ALL AMP Task.xlsx", "VNA", "AMP"),
        (AMOS_DIR / "ALL Customer AMP.xlsx", "CUSTOMER", None),
    ]
    records: list[tuple] = []
    for path, source, sheet in sources:
        if not path.exists():
            continue
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.active
        rows = ws.iter_rows(values_only=True)
        headers = list(next(rows))
        h = cellmap(headers)
        def get(row, *names):
            for name in names:
                idx = h.get(norm(name))
                if idx is not None and idx < len(row):
                    return row[idx]
            return None
        for row in rows:
            task_no = get(row, "Task No./Rev No.", "Jobcard No.", "Taskcard No./Rev")
            if not task_no:
                continue
            records.append((
                source, get(row, "A/C Type"), get(row, "Customer"), str(task_no).strip(), norm(task_no),
                get(row, "Taskcard Title", "Title"), get(row, "Description"), get(row, "Phase"),
                get(row, "Zone"), get(row, "Area"), get(row, "Panel"), get(row, "Special Code"),
                get(row, "Resource Type"), get(row, "Classification"), get(row, "Type Code"), get(row, "Qualification"),
                get(row, "Planning MHR"), get(row, "Information Skill"), get(row, "References"),
                get(row, "Linked Jobcards"), get(row, "Jobcards Title"),
            ))
    with db() as conn:
        conn.executemany(
            """INSERT INTO amp_task(source,aircraft_type,customer,task_no,task_key,title,description,
            phase,zone,area,panel,special_code,resource_type,classification,type_code,qualification,planning_mhr,
            information_skill,requirement_link,linked_jobcards,jobcards_title) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            records,
        )


def import_jobcards() -> None:
    """Import the AMOS Jobcard master into its own read-only MiniDB table."""
    with db() as conn:
        if conn.execute("SELECT COUNT(*) FROM jobcard_master").fetchone()[0]:
            return
    path = AMOS_DIR / "All Jobcard.xlsx"
    if not path.exists():
        return
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    headers = list(next(rows))
    h = cellmap(headers)

    def get(row, name):
        idx = h.get(norm(name))
        return row[idx] if idx is not None and idx < len(row) else None

    records: list[tuple[Any, ...]] = []
    for row in rows:
        jobcard_no = get(row, "Jobcard No.")
        if not jobcard_no:
            continue
        meta_key = get(row, "Meta Key")
        meta_value = get(row, "Meta Value")
        aircraft_type = meta_value if norm(meta_key) == "AIRCRAFT TYPE" else None
        records.append((
            meta_key, meta_value, aircraft_type, str(jobcard_no).strip(), norm(jobcard_no),
            get(row, "Title"), get(row, "Description"), get(row, "Phase"),
            get(row, "Zone"), get(row, "Area"), get(row, "Panel"),
            get(row, "Special Code"), get(row, "Resource Type"),
            get(row, "Classification"), get(row, "Qualification"),
            get(row, "Planning MHR"), get(row, "Information Skill"),
            get(row, "Information MHR"), get(row, "References"),
        ))
    with db() as conn:
        conn.executemany(
            """INSERT INTO jobcard_master(meta_key,meta_value,aircraft_type,jobcard_no,
            jobcard_key,title,description,phase,zone,area,panel,special_code,resource_type,
            classification,qualification,planning_mhr,information_skill,information_mhr,
            requirement_link) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            records,
        )


def backfill_output_columns() -> None:
    """Populate fields added by the approved output workbook in existing demo data."""
    with db() as conn:
        missing_amp = conn.execute(
            "SELECT COUNT(*) FROM amp_task WHERE type_code IS NULL AND jobcards_title IS NULL"
        ).fetchone()[0]
    if missing_amp:
        updates: list[tuple[Any, Any, str, str]] = []
        for path, source, sheet in ((AMOS_DIR / "ALL AMP Task.xlsx", "VNA", "AMP"),
                                    (AMOS_DIR / "ALL Customer AMP.xlsx", "CUSTOMER", None)):
            if not path.exists():
                continue
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb[sheet] if sheet else wb.active
            rows = ws.iter_rows(values_only=True); headers = list(next(rows)); h = cellmap(headers)
            for row in rows:
                def value(name: str):
                    idx = h.get(norm(name)); return row[idx] if idx is not None and idx < len(row) else None
                task_no = value("Task No./Rev No.") or value("Jobcard No.") or value("Taskcard No./Rev")
                if task_no:
                    updates.append((value("Type Code"), value("Jobcards Title"), source, norm(task_no)))
        with db() as conn:
            conn.executemany(
                "UPDATE amp_task SET type_code=?,jobcards_title=? WHERE source=? AND task_key=?",
                updates,
            )
    with db() as conn:
        pending = conn.execute("""SELECT i.id,i.task_no,t.customer,t.aircraft_type
          FROM evaluation_item i JOIN evaluation e ON e.id=i.evaluation_id
          JOIN tally_sheet t ON t.id=e.tally_id
          WHERE i.matched=1 AND i.type_code IS NULL AND i.jobcards_title IS NULL""").fetchall()
        for item in pending:
            amp = lookup_amp(conn, item["task_no"], item["customer"], item["aircraft_type"])
            if amp:
                conn.execute("UPDATE evaluation_item SET type_code=?,jobcards_title=? WHERE id=?",
                             (amp["type_code"], amp["jobcards_title"], item["id"]))


def parse_vietjet(path: Path) -> list[dict[str, Any]]:
    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = ws.iter_rows(values_only=True)
    headers = list(next(rows))
    h = cellmap(headers)
    required = ["WORK STEP / EVENT NUMBER", "TITLE", "EVENT"]
    if any(key not in h for key in required):
        raise ValueError("Không nhận diện được cấu trúc VietJet: thiếu Work Step / Event Number, Title hoặc Event.")
    output = []
    for row in rows:
        task = row[h["WORK STEP / EVENT NUMBER"]]
        if task in (None, ""):
            continue
        def val(name):
            idx = h.get(name)
            return row[idx] if idx is not None and idx < len(row) else None
        output.append({"seq": val("SEQ. NO."), "section": "TALLY SHEET", "task_no": task,
                       "source_title": val("TITLE"), "event_type": val("EVENT"),
                       "rii": val("RII"), "source_note": val("NOTE")})
    return output


def parse_lufthansa(path: Path) -> list[dict[str, Any]]:
    ws = load_workbook(path, read_only=True, data_only=True).active
    matrix = [list(row) for row in ws.iter_rows(values_only=True)]
    output: list[dict[str, Any]] = []
    section = ""
    headers: dict[str, int] = {}
    for row in matrix:
        first = str(row[0]).strip() if row and row[0] not in (None, "") else ""
        if re.match(r"^\d+\.\s", first):
            section = first
            headers = {}
            continue
        rowmap = cellmap(row)
        if "CRI" in rowmap and ("TITEL" in rowmap or "TITLE" in rowmap):
            headers = rowmap
            continue
        if not headers or "CRI" not in headers:
            continue
        cri_idx = headers["CRI"]
        cri = row[cri_idx] if cri_idx < len(row) else None
        if cri in (None, "", "CRI"):
            continue
        def val(*names):
            for name in names:
                idx = headers.get(name)
                if idx is not None and idx < len(row):
                    return row[idx]
            return None
        output.append({"seq": len(output) + 1, "section": section or "OTHER", "task_no": cri,
                       "source_title": val("TITEL", "TITLE", "MS DESCRIPTION"),
                       "event_type": "MS Item" if "MS ITEMS" in section.upper() else section,
                       "rii": None, "source_note": val("MS DESCRIPTION", "DESCRIPTION"),
                       "source_zone": val("ZONE"), "source_ata": val("ATA")})
    if not output:
        raise ValueError("Không nhận diện được các section có cột CRI trong file Lufthansa.")
    return output


def lookup_amp(conn: sqlite3.Connection, task_no: Any, customer: str, aircraft_type: str) -> sqlite3.Row | None:
    key = norm(task_no)
    preferred_source = "CUSTOMER" if customer in {"VA-VJ", "DLH"} else "VNA"
    aircraft_key = canonical_aircraft(aircraft_type)
    return conn.execute(
        """SELECT * FROM amp_task WHERE task_key=?
        ORDER BY CASE WHEN source=? THEN 0 ELSE 1 END,
        CASE WHEN UPPER(COALESCE(customer,''))=UPPER(?) THEN 0 ELSE 1 END,
        CASE WHEN UPPER(COALESCE(aircraft_type,''))=UPPER(?) THEN 0 ELSE 1 END,
        id ASC LIMIT 1""",
        (key, preferred_source, customer, aircraft_key),
    ).fetchone()


WORKFLOW_STEPS = [
    ("intake", "Tiếp nhận Tally Sheet"),
    ("inspect", "Đọc và kiểm tra sơ bộ"),
    ("evaluate", "Chạy đánh giá sơ bộ"),
    ("review", "Review kết quả"),
    ("finalize", "Hoàn thiện và xuất báo cáo"),
]


def ensure_workflow_steps(conn: sqlite3.Connection, tally_id: str) -> None:
    for code, _ in WORKFLOW_STEPS:
        conn.execute(
            """INSERT OR IGNORE INTO workflow_step(tally_id,step_code,status,details_json,updated_at)
            VALUES (?,?,?,'{}',?)""",
            (tally_id, code, "not_started", now()),
        )


def decode_details(value: str | None) -> dict[str, Any]:
    try:
        data = json.loads(value or "{}")
        return data if isinstance(data, dict) else {}
    except json.JSONDecodeError:
        return {}


def set_workflow_step(conn: sqlite3.Connection, tally_id: str, step_code: str,
                      status: str, details: dict[str, Any] | None = None) -> None:
    ensure_workflow_steps(conn, tally_id)
    current = conn.execute(
        "SELECT details_json FROM workflow_step WHERE tally_id=? AND step_code=?",
        (tally_id, step_code),
    ).fetchone()
    payload = decode_details(current["details_json"] if current else None)
    if details is not None:
        payload.update(details)
    completed_at = now() if status == "completed" else None
    conn.execute(
        """UPDATE workflow_step SET status=?,details_json=?,completed_at=?,updated_at=?
        WHERE tally_id=? AND step_code=?""",
        (status, json.dumps(payload, ensure_ascii=False, default=str), completed_at,
         now(), tally_id, step_code),
    )


def workflow_progress(conn: sqlite3.Connection, tally_id: str) -> dict[str, Any]:
    ensure_workflow_steps(conn, tally_id)
    rows = {r["step_code"]: r for r in conn.execute(
        "SELECT * FROM workflow_step WHERE tally_id=?", (tally_id,)
    )}
    steps = []
    for index, (code, label) in enumerate(WORKFLOW_STEPS, 1):
        row = rows[code]
        steps.append({
            "number": index,
            "code": code,
            "label": label,
            "status": row["status"],
            "details": decode_details(row["details_json"]),
            "completed_at": row["completed_at"],
            "updated_at": row["updated_at"],
        })
    completed = sum(step["status"] == "completed" for step in steps)
    current = next((step for step in steps if step["status"] != "completed"), steps[-1])
    return {"steps": steps, "progress": completed * 20, "current_step": current["number"],
            "current_step_code": current["code"], "completed_steps": completed}


def create_tally_record(path: Path, file_name: str, customer: str, aircraft_reg: str,
                        aircraft_type: str, check_type: str, uploader: str,
                        description: str = "") -> str:
    adapter = "VietJet" if customer == "VA-VJ" else "Lufthansa"
    tally_id = f"TS-{datetime.now():%Y%m%d}-{uuid.uuid4().hex[:5].upper()}"
    with db() as conn:
        conn.execute("""INSERT INTO tally_sheet(id,file_name,stored_path,customer,aircraft_reg,aircraft_type,
                     check_type,adapter,uploader,uploaded_at,row_count,status,description)
                     VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                     (tally_id, file_name, str(path), customer, aircraft_reg, aircraft_type,
                      check_type, adapter, uploader, now(), 0, "Đã tiếp nhận", description.strip()))
        ensure_workflow_steps(conn, tally_id)
        set_workflow_step(conn, tally_id, "intake", "completed", {
            "file_name": file_name,
            "description": description.strip(),
            "received_at": now(),
        })
    return tally_id


def inspect_tally_record(tally_id: str) -> dict[str, Any]:
    with db() as conn:
        tally = conn.execute("SELECT * FROM tally_sheet WHERE id=?", (tally_id,)).fetchone()
    if not tally:
        raise HTTPException(404, "Không tìm thấy Tally Sheet")
    path = Path(tally["stored_path"])
    if not path.exists():
        raise HTTPException(404, "Không tìm thấy file Tally Sheet gốc")
    try:
        parsed = parse_vietjet(path) if tally["adapter"] == "VietJet" else parse_lufthansa(path)
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        detected_columns: list[str] = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 60), values_only=True):
            columns = [str(value).strip() for value in row if value not in (None, "")]
            normalized = {norm(value) for value in columns}
            if tally["adapter"] == "VietJet" and "WORK STEP / EVENT NUMBER" in normalized:
                detected_columns = columns
                break
            if tally["adapter"] == "Lufthansa" and "CRI" in normalized:
                detected_columns = columns
                break
        task_keys = [norm(item.get("task_no")) for item in parsed]
        duplicate_count = len(task_keys) - len(set(task_keys))
        warnings = []
        for label, value in (("Số đăng ký tàu bay", tally["aircraft_reg"]),
                             ("Loại tàu bay", tally["aircraft_type"]),
                             ("Check Type", tally["check_type"])):
            if not str(value or "").strip():
                warnings.append(f"Chưa có {label}.")
        if duplicate_count:
            warnings.append(f"Có {duplicate_count} mã task trùng trong Tally Sheet.")
        result = {
            "adapter": tally["adapter"],
            "sheet_names": wb.sheetnames,
            "active_sheet": ws.title,
            "source_rows": ws.max_row,
            "source_columns": ws.max_column,
            "detected_columns": detected_columns,
            "row_count": len(parsed),
            "duplicate_count": duplicate_count,
            "warnings": warnings,
            "confirmed": False,
            "read_at": now(),
            "sample_rows": [{
                "seq": item.get("seq"),
                "task_no": item.get("task_no"),
                "title": item.get("source_title"),
                "section": item.get("section") or item.get("event_type"),
            } for item in parsed[:5]],
        }
    except HTTPException:
        raise
    except Exception as exc:
        with db() as conn:
            conn.execute("UPDATE tally_sheet SET status=?,error_message=? WHERE id=?",
                         ("Lỗi đọc file", str(exc), tally_id))
            set_workflow_step(conn, tally_id, "inspect", "in_progress",
                              {"error": str(exc), "read_at": now()})
        raise HTTPException(400, str(exc)) from exc
    with db() as conn:
        conn.execute("""UPDATE tally_sheet SET row_count=?,status=?,error_message=NULL
                      WHERE id=?""", (len(parsed), "Đã đọc sơ bộ", tally_id))
        set_workflow_step(conn, tally_id, "inspect", "in_progress", result)
    return result


def confirm_tally_inspection(tally_id: str) -> dict[str, Any]:
    with db() as conn:
        row = conn.execute("""SELECT * FROM workflow_step
          WHERE tally_id=? AND step_code='inspect'""", (tally_id,)).fetchone()
        if not row:
            raise HTTPException(400, "Cần chạy đọc sơ bộ trước khi xác nhận")
        details = decode_details(row["details_json"])
        if not details.get("row_count"):
            raise HTTPException(400, "Kết quả đọc sơ bộ chưa hợp lệ")
        details["confirmed"] = True
        details["confirmed_at"] = now()
        set_workflow_step(conn, tally_id, "inspect", "completed", details)
    return details


def evaluate_tally(tally_id: str) -> tuple[str, bool]:
    with db() as conn:
        tally = conn.execute("SELECT * FROM tally_sheet WHERE id=?", (tally_id,)).fetchone()
        if not tally:
            raise HTTPException(404, "Không tìm thấy Tally Sheet")
        existing = conn.execute("SELECT id FROM evaluation WHERE tally_id=?", (tally_id,)).fetchone()
        if existing:
            return existing["id"], True
        inspection = conn.execute("""SELECT details_json FROM workflow_step
          WHERE tally_id=? AND step_code='inspect'""", (tally_id,)).fetchone()
        inspection_details = decode_details(inspection["details_json"]) if inspection else {}
        if not inspection_details.get("row_count"):
            raise HTTPException(400, "Cần chạy đọc sơ bộ trước khi đánh giá")
        if not inspection_details.get("confirmed"):
            raise HTTPException(400, "Cần xác nhận dữ liệu đọc sơ bộ trước khi đánh giá")
    path = Path(tally["stored_path"])
    parsed = parse_vietjet(path) if tally["adapter"] == "VietJet" else parse_lufthansa(path)
    evaluation_id = f"EV-{datetime.now():%Y%m%d}-{uuid.uuid4().hex[:5].upper()}"
    matched = 0
    with db() as conn:
        existing = conn.execute("SELECT id FROM evaluation WHERE tally_id=?", (tally_id,)).fetchone()
        if existing:
            return existing["id"], True
        conn.execute("""INSERT INTO evaluation(id,tally_id,revision,created_at,updated_at)
                     VALUES (?,?,?,?,?)""", (evaluation_id,tally_id,"REV00",now(),now()))
        for item in parsed:
            amp = lookup_amp(conn, item["task_no"], tally["customer"], tally["aircraft_type"])
            if amp:
                matched += 1
                important = 1 if norm(amp["classification"]) == "WATCH ITEM" else 0
                values = (1, amp["source"], amp["title"], amp["description"], amp["phase"],
                          amp["zone"] or item.get("source_zone"), amp["area"], amp["panel"], amp["special_code"],
                          amp["resource_type"], amp["classification"], amp["qualification"],
                          mhr_hours(amp["planning_mhr"]), amp["information_skill"], amp["requirement_link"],
                          amp["linked_jobcards"], amp["type_code"], amp["jobcards_title"],
                          important, "AMOS Watch Item" if important else None, 0)
            else:
                values = (0,None,item.get("source_title"),item.get("source_note"),None,item.get("source_zone"),
                          None,None,None,None,None,None,None,None,None,None,None,None,0,None,1)
            conn.execute("""INSERT INTO evaluation_item(evaluation_id,seq,section,task_no,source_title,event_type,rii,
              source_note,matched,amp_source,title,description,phase,zone,area,panel,special_code,resource_type,
              classification,qualification,planning_mhr,information_skill,requirement_link,linked_jobcards,
              type_code,jobcards_title,important_job,important_source,not_in_database,updated_at)
              VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
              (evaluation_id,item.get("seq"),item.get("section"),str(item["task_no"]),item.get("source_title"),
               item.get("event_type"),item.get("rii"),item.get("source_note"),*values,now()))
        conn.execute("""UPDATE tally_sheet SET matched_count=?,row_count=?,status=? WHERE id=?""",
                     (matched, len(parsed), "Đã đánh giá sơ bộ", tally_id))
        set_workflow_step(conn, tally_id, "evaluate", "completed", {
            "evaluation_id": evaluation_id,
            "total": len(parsed),
            "matched": matched,
            "unmatched": len(parsed) - matched,
            "completed_at": now(),
        })
        set_workflow_step(conn, tally_id, "review", "in_progress", {
            "checklist": {"detail": False, "important": False, "database": False},
            "scope": "not_applicable",
        })
        set_workflow_step(conn, tally_id, "finalize", "not_started", {})
    return evaluation_id, False


def create_evaluation(path: Path, file_name: str, customer: str, aircraft_reg: str,
                      aircraft_type: str, check_type: str, uploader: str,
                      description: str = "") -> str:
    """Legacy all-in-one flow kept for existing callers and sample seeding."""
    tally_id = create_tally_record(path, file_name, customer, aircraft_reg, aircraft_type,
                                   check_type, uploader, description)
    inspect_tally_record(tally_id)
    confirm_tally_inspection(tally_id)
    evaluation_id, _ = evaluate_tally(tally_id)
    return evaluation_id


def seed_samples() -> None:
    with db() as conn:
        if conn.execute("SELECT COUNT(*) FROM tally_sheet").fetchone()[0]:
            return
    samples = [
        ("(Sample) TALLY SHEET WP VJC-A699-24MT2026 REV02.xlsx", "VA-VJ", "VJC-A699", "A320", "24MT"),
        ("(Sample) Workpackage-Summary for D-AIXB_H-25-29102025.xlsx", "DLH", "D-AIXB", "A350", "H-CHECK"),
    ]
    for name, customer, reg, aircraft, check in samples:
        src = SAMPLES_DIR / name
        if src.exists():
            dest = UPLOADS / f"sample-{name}"
            shutil.copy2(src, dest)
            create_evaluation(dest, name, customer, reg, aircraft, check, "Dữ liệu mẫu")


def migrate_workflows() -> None:
    """Backfill workflow state without assuming that legacy reviews were completed."""
    with db() as conn:
        tallies = conn.execute("SELECT * FROM tally_sheet").fetchall()
        for tally in tallies:
            ensure_workflow_steps(conn, tally["id"])
            steps = {r["step_code"]: r for r in conn.execute(
                "SELECT * FROM workflow_step WHERE tally_id=?", (tally["id"],)
            )}
            if steps["intake"]["status"] != "completed":
                set_workflow_step(conn, tally["id"], "intake", "completed", {
                    "file_name": tally["file_name"],
                    "description": tally["description"] or "",
                    "received_at": tally["uploaded_at"],
                    "migrated": True,
                })
            evaluation = conn.execute(
                "SELECT * FROM evaluation WHERE tally_id=?", (tally["id"],)
            ).fetchone()
            if tally["row_count"] and steps["inspect"]["status"] == "not_started":
                set_workflow_step(conn, tally["id"], "inspect", "completed", {
                    "adapter": tally["adapter"],
                    "row_count": tally["row_count"],
                    "warnings": [],
                    "confirmed": True,
                    "migrated": True,
                })
            if evaluation and steps["evaluate"]["status"] != "completed":
                set_workflow_step(conn, tally["id"], "evaluate", "completed", {
                    "evaluation_id": evaluation["id"],
                    "total": tally["row_count"],
                    "matched": tally["matched_count"],
                    "unmatched": max(0, tally["row_count"] - tally["matched_count"]),
                    "migrated": True,
                })
            if evaluation and steps["review"]["status"] == "not_started":
                set_workflow_step(conn, tally["id"], "review", "in_progress", {
                    "checklist": {"detail": False, "important": False, "database": False},
                    "scope": "not_applicable",
                    "migrated": True,
                })
            if evaluation:
                final_status = "completed" if evaluation["status"] in {"Chờ phê duyệt", "Hoàn tất"} else "not_started"
                if steps["finalize"]["status"] != final_status:
                    set_workflow_step(conn, tally["id"], "finalize", final_status, {
                        "evaluation_status": evaluation["status"],
                        "migrated": True,
                    })


@app.on_event("startup")
def startup():
    init_db()
    import_amos()
    import_jobcards()
    backfill_output_columns()
    if os.environ.get("WP_SKIP_SEED_SAMPLES") != "1":
        seed_samples()
    migrate_workflows()


@app.get("/")
def index():
    return FileResponse(ROOT / "templates" / "index.html")


@app.get("/api/dashboard")
def dashboard():
    with db() as conn:
        totals = conn.execute("""SELECT COUNT(*) tally_count, COALESCE(SUM(row_count),0) item_count,
          COALESCE(SUM(matched_count),0) matched_count FROM tally_sheet""").fetchone()
        records = conn.execute("""SELECT t.id tally_id,t.file_name,t.description,t.customer,
          t.aircraft_reg,t.aircraft_type,t.check_type,t.adapter,t.row_count,t.matched_count,
          t.uploaded_at,t.status tally_status,e.id evaluation_id,e.status evaluation_status,
          e.created_at,e.updated_at FROM tally_sheet t LEFT JOIN evaluation e ON e.tally_id=t.id
          ORDER BY t.uploaded_at DESC""").fetchall()
        workflows = []
        for record in records:
            item = dict(record)
            item.update(workflow_progress(conn, record["tally_id"]))
            item["status"] = record["evaluation_status"] or record["tally_status"]
            workflows.append(item)
        evaluations = [item for item in workflows if item["evaluation_id"]]
        amp_count = conn.execute("SELECT COUNT(*) FROM amp_task").fetchone()[0]
    return {"totals": dict(totals), "amp_count": amp_count,
            "workflows": workflows, "evaluations": evaluations}


@app.get("/api/data/amos")
def browse_amos(q: str = "", source: str = "all", aircraft_type: str = "all",
                customer: str = "all", classification: str = "all", page: int = 1,
                page_size: int = 40):
    page = max(1, page); page_size = min(max(10, page_size), 100)
    where = ["1=1"]; params: list[Any] = []
    if q.strip():
        where.append("(task_no LIKE ? OR title LIKE ? OR description LIKE ? OR special_code LIKE ?)")
        params.extend([f"%{q.strip()}%"] * 4)
    for column, value in (("source", source), ("aircraft_type", aircraft_type),
                          ("customer", customer), ("classification", classification)):
        if value != "all":
            where.append(f"COALESCE({column},'')=?"); params.append(value)
    predicate = " AND ".join(where)
    with db() as conn:
        total = conn.execute(f"SELECT COUNT(*) FROM amp_task WHERE {predicate}", params).fetchone()[0]
        rows = [dict(r) for r in conn.execute(f"""SELECT id,source,aircraft_type,customer,task_no,title,
          phase,zone,special_code,classification,planning_mhr,qualification,description,area,panel,
          resource_type,information_skill,requirement_link,linked_jobcards FROM amp_task
          WHERE {predicate} ORDER BY task_no LIMIT ? OFFSET ?""",
          (*params, page_size, (page - 1) * page_size))]
        def distinct(column: str, limit: int = 80):
            return [r[0] for r in conn.execute(f"""SELECT DISTINCT {column} FROM amp_task
              WHERE {column} IS NOT NULL AND TRIM({column})<>'' ORDER BY {column} LIMIT ?""", (limit,))]
        filters = {"sources": distinct("source"), "aircraft_types": distinct("aircraft_type"),
                   "customers": distinct("customer"), "classifications": distinct("classification")}
        stats = dict(conn.execute("""SELECT COUNT(*) total,
          SUM(CASE WHEN source='VNA' THEN 1 ELSE 0 END) vna,
          SUM(CASE WHEN source='CUSTOMER' THEN 1 ELSE 0 END) customer,
          SUM(CASE WHEN UPPER(COALESCE(classification,''))='WATCH ITEM' THEN 1 ELSE 0 END) watch_items
          FROM amp_task""").fetchone())
    return {"items": rows, "total": total, "page": page, "page_size": page_size,
            "filters": filters, "stats": stats}


@app.get("/api/data/jobcards")
def browse_jobcards(q: str = "", aircraft_type: str = "all",
                    classification: str = "all", page: int = 1, page_size: int = 40):
    page = max(1, page); page_size = min(max(10, page_size), 100)
    where = ["1=1"]; params: list[Any] = []
    if q.strip():
        where.append("(jobcard_no LIKE ? OR title LIKE ? OR description LIKE ? OR special_code LIKE ?)")
        params.extend([f"%{q.strip()}%"] * 4)
    for column, value in (("aircraft_type", aircraft_type), ("classification", classification)):
        if value != "all":
            where.append(f"COALESCE({column},'')=?"); params.append(value)
    predicate = " AND ".join(where)
    with db() as conn:
        total = conn.execute(f"SELECT COUNT(*) FROM jobcard_master WHERE {predicate}", params).fetchone()[0]
        rows = [dict(r) for r in conn.execute(f"""SELECT id,'JOBCARD' source,aircraft_type,
          NULL customer,jobcard_no task_no,title,description,phase,zone,area,panel,special_code,
          resource_type,classification,qualification,planning_mhr,information_skill,
          information_mhr,requirement_link FROM jobcard_master
          WHERE {predicate} ORDER BY jobcard_no LIMIT ? OFFSET ?""",
          (*params, page_size, (page - 1) * page_size))]

        def distinct(column: str, limit: int = 80):
            return [r[0] for r in conn.execute(f"""SELECT DISTINCT {column} FROM jobcard_master
              WHERE {column} IS NOT NULL AND TRIM({column})<>'' ORDER BY {column} LIMIT ?""", (limit,))]

        filters = {"aircraft_types": distinct("aircraft_type"),
                   "classifications": distinct("classification")}
        stats = dict(conn.execute("""SELECT COUNT(*) total,
          SUM(CASE WHEN UPPER(COALESCE(classification,''))='WATCH ITEM' THEN 1 ELSE 0 END) watch_items
          FROM jobcard_master""").fetchone())
    return {"items": rows, "total": total, "page": page, "page_size": page_size,
            "filters": filters, "stats": stats}


@app.get("/api/data/external")
def browse_external(q: str = "", category: str = "all", status: str = "all",
                    page: int = 1, page_size: int = 40):
    page = max(1, page); page_size = min(max(10, page_size), 100)
    where = ["1=1"]; params: list[Any] = []
    if q.strip():
        where.append("(code LIKE ? OR canonical_value LIKE ? OR description LIKE ? OR reason LIKE ?)")
        params.extend([f"%{q.strip()}%"] * 4)
    if category != "all": where.append("category=?"); params.append(category)
    if status != "all": where.append("status=?"); params.append(status)
    predicate = " AND ".join(where)
    with db() as conn:
        total = conn.execute(f"SELECT COUNT(*) FROM external_data WHERE {predicate}", params).fetchone()[0]
        rows = [dict(r) for r in conn.execute(f"""SELECT * FROM external_data WHERE {predicate}
          ORDER BY category,code LIMIT ? OFFSET ?""", (*params,page_size,(page-1)*page_size))]
        categories = [r[0] for r in conn.execute("SELECT DISTINCT category FROM external_data ORDER BY category")]
        statuses = [r[0] for r in conn.execute("SELECT DISTINCT status FROM external_data ORDER BY status")]
        category_counts = [dict(r) for r in conn.execute("""SELECT category,COUNT(*) record_count
          FROM external_data GROUP BY category ORDER BY category""")]
    return {"items":rows,"total":total,"page":page,"page_size":page_size,
            "filters":{"categories":categories,"statuses":statuses},"category_counts":category_counts}


async def store_upload(file: UploadFile) -> Path:
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(400, "Chỉ hỗ trợ file Excel .xlsx")
    safe_name = re.sub(r"[^\w.() -]", "_", Path(file.filename).name)
    dest = UPLOADS / f"{datetime.now():%Y%m%d%H%M%S}-{uuid.uuid4().hex[:5]}-{safe_name}"
    content = await file.read()
    if len(content) > 30 * 1024 * 1024:
        raise HTTPException(400, "File vượt quá giới hạn 30 MB")
    dest.write_bytes(content)
    return dest


@app.post("/api/tallies")
async def create_tally(file: UploadFile = File(...), customer: str = Form(...),
                       aircraft_reg: str = Form(""), aircraft_type: str = Form(""),
                       check_type: str = Form(""), description: str = Form(""),
                       uploader: str = Form("Demo User")):
    dest = await store_upload(file)
    try:
        tally_id = create_tally_record(dest, file.filename or dest.name, customer, aircraft_reg,
                                       aircraft_type, check_type, uploader, description)
    except Exception as exc:
        dest.unlink(missing_ok=True)
        raise HTTPException(400, str(exc)) from exc
    return {"tally_id": tally_id, "next_step": 2}


@app.post("/api/tallies/{tally_id}/inspect")
def inspect_tally(tally_id: str):
    return {"inspection": inspect_tally_record(tally_id)}


@app.post("/api/tallies/{tally_id}/inspection/confirm")
def confirm_inspection(tally_id: str):
    return {"inspection": confirm_tally_inspection(tally_id)}


@app.post("/api/tallies/{tally_id}/evaluate")
def run_preliminary_evaluation(tally_id: str):
    evaluation_id, reused = evaluate_tally(tally_id)
    return {"evaluation_id": evaluation_id, "reused": reused}


@app.get("/api/workflows/{tally_id}")
def get_workflow(tally_id: str):
    with db() as conn:
        tally = conn.execute("SELECT * FROM tally_sheet WHERE id=?", (tally_id,)).fetchone()
        if not tally:
            raise HTTPException(404, "Không tìm thấy hồ sơ")
        evaluation = conn.execute("SELECT * FROM evaluation WHERE tally_id=?", (tally_id,)).fetchone()
        summary = None
        if evaluation:
            summary = dict(conn.execute("""SELECT COUNT(*) total,COALESCE(SUM(matched),0) matched,
              COALESCE(SUM(not_in_database),0) unmatched,COALESCE(SUM(important_job),0) important,
              ROUND(COALESCE(SUM(COALESCE(reviewed_mhr,planning_mhr)),0),2) total_mhr
              FROM evaluation_item WHERE evaluation_id=?""", (evaluation["id"],)).fetchone())
        progress = workflow_progress(conn, tally_id)
    return {"tally": dict(tally), "evaluation": dict(evaluation) if evaluation else None,
            "summary": summary, **progress}


class ReviewChecklistUpdate(BaseModel):
    detail: bool = False
    important: bool = False
    database: bool = False


@app.patch("/api/workflows/{tally_id}/review-checklist")
def update_review_checklist(tally_id: str, payload: ReviewChecklistUpdate):
    checklist = payload.model_dump()
    with db() as conn:
        evaluation = conn.execute("SELECT id FROM evaluation WHERE tally_id=?", (tally_id,)).fetchone()
        if not evaluation:
            raise HTTPException(400, "Cần chạy đánh giá sơ bộ trước khi review")
        complete = all(checklist.values())
        set_workflow_step(conn, tally_id, "review", "completed" if complete else "in_progress", {
            "checklist": checklist,
            "scope": "not_applicable",
            "confirmed_at": now() if complete else None,
        })
        progress = workflow_progress(conn, tally_id)
    return {"ok": True, "review_complete": complete, **progress}


@app.post("/api/upload")
async def upload(file: UploadFile = File(...), customer: str = Form(...), aircraft_reg: str = Form(""),
                 aircraft_type: str = Form(""), check_type: str = Form(""),
                 description: str = Form(""), uploader: str = Form("Demo User")):
    """Legacy all-in-one endpoint retained for compatibility."""
    dest = await store_upload(file)
    try:
        evaluation_id = create_evaluation(dest, file.filename or dest.name, customer, aircraft_reg,
                                          aircraft_type, check_type, uploader, description)
    except Exception as exc:
        dest.unlink(missing_ok=True)
        raise HTTPException(400, str(exc)) from exc
    return {"evaluation_id": evaluation_id}


@app.get("/api/evaluations/{evaluation_id}")
def evaluation_detail(evaluation_id: str, q: str = "", status: str = "all", page: int = 1,
                      page_size: int = 50, sort_by: str = "seq", sort_dir: str = "asc"):
    page = max(1, page); page_size = min(max(10,page_size),200)
    where = ["i.evaluation_id=?"]; params: list[Any] = [evaluation_id]
    if q:
        where.append("(i.task_no LIKE ? OR i.title LIKE ? OR i.description LIKE ?)")
        params.extend([f"%{q}%"] * 3)
    if status == "matched": where.append("i.matched=1")
    if status == "unmatched": where.append("i.matched=0")
    if status == "important": where.append("i.important_job=1")
    predicate = " AND ".join(where)
    with db() as conn:
        header = conn.execute("""SELECT e.*,t.file_name,t.description,t.customer,t.aircraft_reg,t.aircraft_type,t.check_type,
          t.adapter,t.version,t.uploaded_at,t.uploader FROM evaluation e JOIN tally_sheet t ON t.id=e.tally_id WHERE e.id=?""",
          (evaluation_id,)).fetchone()
        if not header: raise HTTPException(404, "Không tìm thấy lần đánh giá")
        total = conn.execute(f"SELECT COUNT(*) FROM evaluation_item i WHERE {predicate}", params).fetchone()[0]
        sortable = {"task_no", "title", "description", "phase", "zone", "area", "panel", "special_code",
                    "resource_type", "classification", "type_code", "qualification", "planning_mhr",
                    "information_skill", "requirement_link", "linked_jobcards", "jobcards_title",
                    "important_job", "not_in_database", "seq"}
        order_col = sort_by if sort_by in sortable else "seq"
        direction = "DESC" if sort_dir.lower() == "desc" else "ASC"
        items = [dict(r) for r in conn.execute(f"SELECT i.* FROM evaluation_item i WHERE {predicate} ORDER BY i.{order_col} {direction},i.id ASC LIMIT ? OFFSET ?",
                 (*params,page_size,(page-1)*page_size))]
        summary = dict(conn.execute("""SELECT COUNT(*) total, SUM(matched) matched, SUM(not_in_database) unmatched,
          SUM(important_job) important, ROUND(SUM(COALESCE(reviewed_mhr,planning_mhr)),2) total_mhr
          FROM evaluation_item WHERE evaluation_id=?""", (evaluation_id,)).fetchone())
        special = [dict(r) for r in conn.execute("""SELECT COALESCE(NULLIF(special_code,''),'Chưa phân loại') special_code,
          COUNT(*) task_count,ROUND(SUM(COALESCE(reviewed_mhr,planning_mhr)),2) mhr FROM evaluation_item
          WHERE evaluation_id=? GROUP BY 1 ORDER BY mhr DESC LIMIT 12""", (evaluation_id,))]
    return {"header":dict(header),"summary":summary,"items":items,"special_codes":special,"filtered_total":total,"page":page,"page_size":page_size}


class ItemUpdate(BaseModel):
    reviewed_mhr: float | None = None
    important_job: bool | None = None
    review_note: str | None = None
    reason: str
    changed_by: str = "Demo User"


@app.patch("/api/items/{item_id}")
def update_item(item_id: int, payload: ItemUpdate):
    if not payload.reason.strip(): raise HTTPException(400,"Bắt buộc nhập lý do thay đổi")
    with db() as conn:
        row = conn.execute("SELECT * FROM evaluation_item WHERE id=?",(item_id,)).fetchone()
        if not row: raise HTTPException(404,"Không tìm thấy task")
        changes = {k:v for k,v in {"reviewed_mhr":payload.reviewed_mhr,"important_job":int(payload.important_job) if payload.important_job is not None else None,"review_note":payload.review_note}.items() if v is not None}
        for field,value in changes.items():
            if str(row[field] or "") != str(value):
                conn.execute("INSERT INTO change_log(evaluation_id,item_id,field_name,old_value,new_value,reason,changed_by,changed_at) VALUES (?,?,?,?,?,?,?,?)",
                             (row["evaluation_id"],item_id,field,row[field],value,payload.reason,payload.changed_by,now()))
                conn.execute(f"UPDATE evaluation_item SET {field}=?,updated_at=? WHERE id=?",(value,now(),item_id))
        conn.execute("UPDATE evaluation SET updated_at=? WHERE id=?",(now(),row["evaluation_id"]))
    return {"ok":True}


class EvaluationUpdate(BaseModel):
    tat_review: str | None = None
    assumptions: str | None = None
    notes: str | None = None
    prepared_by: str | None = None
    approved_by: str | None = None
    status: str | None = None


@app.patch("/api/evaluations/{evaluation_id}")
def update_evaluation(evaluation_id: str, payload: EvaluationUpdate):
    values = payload.model_dump(exclude_none=True)
    if payload.status and payload.status not in {"Đang review", "Chờ phê duyệt", "Hoàn tất"}:
        raise HTTPException(400, "Trạng thái hồ sơ không hợp lệ")
    if values:
        fields = ",".join(f"{key}=?" for key in values)
        with db() as conn:
            evaluation = conn.execute("SELECT tally_id FROM evaluation WHERE id=?", (evaluation_id,)).fetchone()
            if not evaluation:
                raise HTTPException(404, "Không tìm thấy lần đánh giá")
            conn.execute(f"UPDATE evaluation SET {fields},updated_at=? WHERE id=?",(*values.values(),now(),evaluation_id))
            status = payload.status or conn.execute(
                "SELECT status FROM evaluation WHERE id=?", (evaluation_id,)
            ).fetchone()["status"]
            final_status = "completed" if status in {"Chờ phê duyệt", "Hoàn tất"} else "in_progress"
            set_workflow_step(conn, evaluation["tally_id"], "finalize", final_status, {
                "evaluation_status": status,
                "saved_at": now(),
            })
    return {"ok":True}


@app.get("/api/evaluations/{evaluation_id}/report/{kind}")
def report(evaluation_id: str, kind: str):
    conditions = {"important":"important_job=1","database":"not_in_database=1","scope":"0"}
    if kind not in conditions: raise HTTPException(404,"Báo cáo không tồn tại")
    with db() as conn:
        rows = [dict(r) for r in conn.execute(f"SELECT * FROM evaluation_item WHERE evaluation_id=? AND {conditions[kind]} ORDER BY id",(evaluation_id,))]
    return {"items":rows,"available":kind != "scope","message":"Chưa có dữ liệu Task–Check/Linked Check để đánh giá tự động." if kind=="scope" else None}


def style_sheet(ws):
    blue = "003B71"; gold = "C89B3C"
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=blue); cell.font=Font(color="FFFFFF",bold=True)
        cell.alignment=Alignment(wrap_text=True,vertical="center")
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    for col in range(1,ws.max_column+1):
        width=min(42,max(12,max(len(str(ws.cell(r,col).value or "")) for r in range(1,min(ws.max_row,80)+1))+2))
        ws.column_dimensions[get_column_letter(col)].width=width
    ws.sheet_properties.tabColor=gold


@app.get("/api/evaluations/{evaluation_id}/export")
def export_excel(evaluation_id: str):
    with db() as conn:
        header=conn.execute("""SELECT e.*,t.* FROM evaluation e JOIN tally_sheet t ON t.id=e.tally_id WHERE e.id=?""",(evaluation_id,)).fetchone()
        if not header: raise HTTPException(404,"Không tìm thấy lần đánh giá")
        items=[dict(r) for r in conn.execute("SELECT * FROM evaluation_item WHERE evaluation_id=? ORDER BY id",(evaluation_id,))]
        progress = workflow_progress(conn, header["tally_id"])
    is_draft = progress["progress"] < 100 or header["status"] not in {"Chờ phê duyệt", "Hoàn tất"}
    wb=Workbook(); wb.remove(wb.active)
    columns=[("Task No.","task_no"),("Taskcard Title","title"),("Description","description"),("Phase","phase"),("Zone","zone"),("Area","area"),("Panel","panel"),("Special Code","special_code"),("Resource Type","resource_type"),("Classification","classification"),("Type Code","type_code"),("Qualification","qualification"),("Planning MHR","planning_mhr"),("Information Skill","information_skill"),("Requirement Link","requirement_link"),("Linked Jobcards","linked_jobcards"),("Jobcards Title","jobcards_title"),("Important Job","important_job"),("Not included in Database","not_in_database")]
    for title, subset in [("Detailed Task List",items),("Important Jobs",[x for x in items if x["important_job"]]),("Task out of Check Type Scope",[]) ,("Database to be Updated",[x for x in items if x["not_in_database"]])]:
        ws=wb.create_sheet(title); ws.append([c[0] for c in columns])
        for item in subset:
            row = []
            for _, key in columns:
                if key in {"important_job", "not_in_database"}:
                    row.append("X" if item[key] else None)
                elif key == "planning_mhr" and not item["matched"]:
                    row.append(None)
                else:
                    row.append(item[key])
            ws.append(row)
        if title=="Task out of Check Type Scope": ws.append(["Chưa có dữ liệu Task–Check/Linked Check để đánh giá tự động."])
        style_sheet(ws)
    ws=wb.create_sheet("FINAL EVALUATION",0)
    ws.append(["BẢN NHÁP - FINAL EVALUATION" if is_draft else "FINAL EVALUATION",""]); ws.append(["Thông tin","Giá trị"])
    fields=[("Trạng thái báo cáo","BẢN NHÁP" if is_draft else header["status"]),("Khách hàng",header["customer"]),("Mô tả hồ sơ",header["description"]),("Số đăng ký",header["aircraft_reg"]),("Loại tàu bay",header["aircraft_type"]),("Check Type",header["check_type"]),("Workpack revision",header["revision"]),("TAT review",header["tat_review"]),("Tổng task",len(items)),("Task khớp",sum(x["matched"] for x in items)),("Task thiếu DB",sum(x["not_in_database"] for x in items)),("Important Job",sum(x["important_job"] for x in items)),("Tổng MHR",round(sum((x["reviewed_mhr"] if x["reviewed_mhr"] is not None else x["planning_mhr"]) or 0 for x in items),2)),("Giả định",header["assumptions"]),("Ghi chú",header["notes"]),("Người lập",header["prepared_by"]),("Người phê duyệt",header["approved_by"])]
    for row in fields: ws.append(row)
    ws["A1"].fill=PatternFill("solid",fgColor="B63E3E" if is_draft else "003B71"); ws["A1"].font=Font(color="FFFFFF",bold=True,size=16); ws.merge_cells("A1:B1"); ws.column_dimensions["A"].width=28; ws.column_dimensions["B"].width=60
    stream=io.BytesIO(); wb.save(stream); stream.seek(0)
    name=f"{'DRAFT_' if is_draft else ''}VAECO_WP_Evaluation_{evaluation_id}.xlsx"
    return StreamingResponse(stream,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":f'attachment; filename="{name}"'})


@app.get("/api/tally/{tally_id}/download")
def download_tally(tally_id: str):
    with db() as conn: row=conn.execute("SELECT * FROM tally_sheet WHERE id=?",(tally_id,)).fetchone()
    if not row or not Path(row["stored_path"]).exists(): raise HTTPException(404,"Không tìm thấy file gốc")
    return FileResponse(row["stored_path"],filename=row["file_name"])


@app.get("/api/evaluations/{evaluation_id}/history")
def history(evaluation_id: str):
    with db() as conn: rows=[dict(r) for r in conn.execute("SELECT * FROM change_log WHERE evaluation_id=? ORDER BY changed_at DESC",(evaluation_id,))]
    return {"items":rows}
