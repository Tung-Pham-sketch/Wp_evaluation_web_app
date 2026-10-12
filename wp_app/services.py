# =============================================================================
# services.py – WORKFLOW, EVALUATION, STARTUP MIGRATIONS, DELETE
# -----------------------------------------------------------------------------
# Files
#   portable_stored_path() / resolve_stored_path()
#       upload paths stored relative to data/, so a copied data folder still
#       finds its files on another PC.
#
# 5-step workflow (WORKFLOW_STEPS: intake, inspect, evaluate, review, finalize)
#   create_tally_record()      step 1 – save tally_sheet, adapter from the form
#   inspect_tally_record()     step 2 – parse, count rows, duplicates, structure
#   confirm_tally_inspection() step 2 – user confirms the input
#   evaluate_tally()           step 3 – create evaluation + items (once per tally)
#   set_workflow_step() / workflow_progress()  status per step, 20 % each.
#
# Evaluation  –  populate_evaluation_items()   *** core of step 3 ***
#   for every parsed row:
#     prepare_row() -> not eligible           -> NOT_APPLICABLE
#     layer 1: mpd_matching.match_rows()      -> layer1_fields():
#              AUTO_MATCH, or REVIEW_REQUIRED if a review note exists;
#              MHR and AMOS values are filled in.
#     no layer-1 result: suggestion_decision() = lookup_amp(include_exact=False)
#              -> layer 2 (prefix) / layer 3 (Title) REVIEW_REQUIRED suggestion,
#                 matched = 0, NO MHR, only suggested_* columns
#              -> nothing: NOT_FOUND, not_in_database = 1
#
# lookup_amp(..., include_exact=True)   (fork function, extended)
#   fork's 3-layer lookup. include_exact=False skips its exact layer, which is
#   how the merged flow uses it. Customer-first scope comes from
#   matching_config.CUSTOMER_FIRST (method names e.g. CR_DLH_A350).
#
# Startup migrations (called from main.py)
#   migrate_upload_paths, migrate_legacy_matching_results (re-match when
#   MATCH_LOGIC_VERSION increased and no change_log / amp_change exists),
#   migrate_legacy_status_labels (old Vietnamese labels -> codes),
#   migrate_workflows, seed_samples.
#
# delete_tally()  removes tally, evaluation, items, change_log, workflow and
#   the uploaded file. amp_change and MiniDB records are kept on purpose.
# =============================================================================

from __future__ import annotations

import json
import shutil
import sqlite3
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any

from fastapi import HTTPException
from openpyxl import load_workbook

from .adapters import adapter_for_upload, inspect_tally_structure, parse_tally
from .config import DATA_DIR, MATCH_LOGIC_VERSION, SAMPLES_DIR, UPLOADS_DIR
from .database import db, now
from .matching import (
    AUTO_MATCH,
    NOT_APPLICABLE,
    NOT_FOUND,
    PREFIX_REVIEW_THRESHOLD,
    REVIEW_REQUIRED,
    TITLE_REVIEW_THRESHOLD,
    MatchDecision,
    best_scored_candidate,
    canonical_aircraft,
    canonical_customer,
    resolve_candidates,
    taskcard_core,
)
from .mpd_matching import (
    STEP_METHOD,
    combined_note,
    is_customer_first,
    match_rows,
    prepare_row,
    tally_title_text,
)
from .utils import mhr_hours, normalize


WORKFLOW_STEPS = [
    ("intake", "Tiếp nhận Tally Sheet"),
    ("inspect", "Đọc và kiểm tra sơ bộ"),
    ("evaluate", "Chạy đánh giá sơ bộ"),
    ("review", "Review kết quả"),
    ("finalize", "Hoàn thiện và xuất báo cáo"),
]


def portable_stored_path(path: Path) -> str:
    """Store runtime files relative to DATA_DIR so a copied data folder stays portable."""
    resolved = path.resolve()
    try:
        return str(resolved.relative_to(DATA_DIR))
    except ValueError:
        return str(resolved)


def resolve_stored_path(stored_path: str, file_name: str | None = None) -> Path:
    """Resolve current and legacy upload paths after the project moves to another machine."""
    raw = Path(stored_path)
    candidates = [raw if raw.is_absolute() else DATA_DIR / raw]
    candidates.append(UPLOADS_DIR / raw.name)
    if file_name:
        candidates.extend((UPLOADS_DIR / f"sample-{file_name}", SAMPLES_DIR / file_name))
    for candidate in candidates:
        if candidate.exists():
            return candidate.resolve()
    return candidates[0]


def migrate_upload_paths() -> None:
    """Normalize existing upload paths without changing missing legacy references."""
    with db() as connection:
        rows = connection.execute("SELECT id,file_name,stored_path FROM tally_sheet").fetchall()
        for row in rows:
            resolved = resolve_stored_path(row["stored_path"], row["file_name"])
            if resolved.exists():
                connection.execute(
                    "UPDATE tally_sheet SET stored_path=? WHERE id=?",
                    (portable_stored_path(resolved), row["id"]),
                )


def _match_decision(resolution, method: str) -> MatchDecision:
    is_candidate = resolution.status in {AUTO_MATCH, REVIEW_REQUIRED}
    note = resolution.note
    if resolution.status == AUTO_MATCH:
        note = "Lớp 1 - khớp 100% theo khóa và phạm vi nguồn/customer/model đã thống nhất."
    elif resolution.status == REVIEW_REQUIRED and note:
        note = f"Lớp 1 - mã khớp 100% nhưng dữ liệu candidate xung đột. {note}"
    return MatchDecision(
        status=resolution.status,
        method=method,
        record=resolution.record,
        candidate_count=resolution.candidate_count,
        note=note,
        score=100.0 if is_candidate else None,
        layer=1 if is_candidate else None,
    )


def lookup_amp(connection: sqlite3.Connection, task_no: Any, customer: str,
               aircraft_type: str, source_title: Any = None, *,
               include_exact: bool = True) -> MatchDecision:
    """Apply the approved source route: VNA Taskcard first, then Customer CR.

    Merged: the evaluation runs the MPD-key flow (mpd_matching.match_rows) as layer 1
    and calls this function with include_exact=False, so only layers 2 (prefix)
    and 3 (Title) are used, as suggestions without MHR.
    Customers searched in Customer AMP first come from matching_config.CUSTOMER_FIRST
    (DLH + A350 by default) instead of being written in the code.
    """
    key = normalize(task_no)
    core_key = taskcard_core(task_no)
    customer_key = canonical_customer(customer)
    aircraft_key = canonical_aircraft(aircraft_type)
    customer_pool: list[sqlite3.Row] | None = None
    vna_pool: list[sqlite3.Row] | None = None

    def all_customer_rows() -> list[sqlite3.Row]:
        nonlocal customer_pool
        if customer_pool is None:
            rows = connection.execute(
                "SELECT * FROM amp_task WHERE source='CUSTOMER' ORDER BY id"
            ).fetchall()
            customer_pool = [
                row for row in rows if canonical_customer(row["customer"]) == customer_key
            ]
        return customer_pool

    def exact_customer_rows() -> list[sqlite3.Row]:
        rows = connection.execute(
            "SELECT * FROM amp_task WHERE source='CUSTOMER' AND taskcard_key=? ORDER BY id",
            (key,),
        ).fetchall()
        return [row for row in rows if canonical_customer(row["customer"]) == customer_key]

    def all_vna_rows() -> list[sqlite3.Row]:
        nonlocal vna_pool
        if vna_pool is None:
            vna_pool = connection.execute(
                "SELECT * FROM amp_task WHERE source='VNA' ORDER BY id"
            ).fetchall()
        return vna_pool

    def review_decision(scored, method: str, layer: int, basis: str) -> MatchDecision:
        tied = f" Có {scored.tied_count} candidate đồng hạng cao nhất." if scored.tied_count > 1 else ""
        note = (
            f"Lớp {layer} - {basis}: tỷ lệ {scored.score:.1f}%. "
            f"Có {scored.candidate_count} candidate đạt ngưỡng.{tied} Cần người dùng review."
        )
        return MatchDecision(
            REVIEW_REQUIRED, method, scored.record, scored.candidate_count,
            note, scored.score, layer,
        )

    def prefix_candidate(rows, method: str, requested_aircraft: str) -> MatchDecision | None:
        scored = best_scored_candidate(
            rows, requested_aircraft, core_key, kind="prefix", threshold=PREFIX_REVIEW_THRESHOLD
        )
        return review_decision(scored, method, 2, "khớp tiền tố mã") if scored else None

    def title_candidate(rows, method: str, requested_aircraft: str) -> MatchDecision | None:
        scored = best_scored_candidate(
            rows, requested_aircraft, source_title, kind="title", threshold=TITLE_REVIEW_THRESHOLD
        )
        return review_decision(scored, method, 3, "khớp nội dung Title") if scored else None

    if is_customer_first(customer_key, aircraft_key):
        scope = f"{customer_key}_{aircraft_key}"
        if include_exact:
            exact = resolve_candidates(exact_customer_rows(), aircraft_key)
            if exact.status != NOT_FOUND:
                return _match_decision(exact, f"CR_{scope}")
        prefix = prefix_candidate(all_customer_rows(), f"PREFIX_CR_{scope}", aircraft_key)
        if prefix:
            return prefix
        title = title_candidate(all_customer_rows(), f"TITLE_CR_{scope}", aircraft_key)
        if title:
            return title
        return MatchDecision(NOT_FOUND, f"NO_CANDIDATE_{scope}", None, 0)

    if include_exact:
        vna_rows = connection.execute(
            "SELECT * FROM amp_task WHERE source='VNA' AND taskcard_core_key=? ORDER BY id",
            (core_key,),
        ).fetchall()
        vna_result = resolve_candidates(vna_rows, aircraft_key)
        if vna_result.status != NOT_FOUND:
            return _match_decision(vna_result, "VNA_TASKCARD")
        customer_result = resolve_candidates(exact_customer_rows(), aircraft_key)
        if customer_result.status != NOT_FOUND:
            return _match_decision(customer_result, "CR_FALLBACK")
    prefix = prefix_candidate(all_vna_rows(), "PREFIX_VNA", aircraft_key)
    if prefix:
        return prefix
    prefix = prefix_candidate(all_customer_rows(), "PREFIX_CR", aircraft_key)
    if prefix:
        return prefix
    title = title_candidate(all_vna_rows(), "TITLE_VNA", aircraft_key)
    if title:
        return title
    title = title_candidate(all_customer_rows(), "TITLE_CR", aircraft_key)
    if title:
        return title
    return MatchDecision(NOT_FOUND, "NO_CANDIDATE", None, 0)


def ensure_workflow_steps(connection: sqlite3.Connection, tally_id: str) -> None:
    for code, _ in WORKFLOW_STEPS:
        connection.execute(
            """INSERT OR IGNORE INTO workflow_step(tally_id,step_code,status,details_json,updated_at)
            VALUES (?,?,?,'{}',?)""",
            (tally_id, code, "not_started", now()),
        )


def decode_details(value: str | None) -> dict[str, Any]:
    try:
        data = json.loads(value or "{}")
        return data if isinstance(data, dict) else {}
    except json.JSONDecodeError:
        return {}


def set_workflow_step(connection: sqlite3.Connection, tally_id: str, step_code: str,
                      status: str, details: dict[str, Any] | None = None) -> None:
    ensure_workflow_steps(connection, tally_id)
    current = connection.execute(
        "SELECT details_json FROM workflow_step WHERE tally_id=? AND step_code=?",
        (tally_id, step_code),
    ).fetchone()
    payload = decode_details(current["details_json"] if current else None)
    if details is not None:
        payload.update(details)
    completed_at = now() if status == "completed" else None
    connection.execute(
        """UPDATE workflow_step SET status=?,details_json=?,completed_at=?,updated_at=?
        WHERE tally_id=? AND step_code=?""",
        (
            status,
            json.dumps(payload, ensure_ascii=False, default=str),
            completed_at,
            now(),
            tally_id,
            step_code,
        ),
    )


def workflow_progress(connection: sqlite3.Connection, tally_id: str) -> dict[str, Any]:
    ensure_workflow_steps(connection, tally_id)
    rows = {
        row["step_code"]: row
        for row in connection.execute(
            "SELECT * FROM workflow_step WHERE tally_id=?", (tally_id,)
        )
    }
    steps = []
    for index, (code, label) in enumerate(WORKFLOW_STEPS, 1):
        row = rows[code]
        steps.append({
            "number": index,
            "code": code,
            "label": label,
            "status": row["status"],
            "details": decode_details(row["details_json"]),
            "completed_at": row["completed_at"],
            "updated_at": row["updated_at"],
        })
    completed = sum(step["status"] == "completed" for step in steps)
    current = next((step for step in steps if step["status"] != "completed"), steps[-1])
    return {
        "steps": steps,
        "progress": completed * 20,
        "current_step": current["number"],
        "current_step_code": current["code"],
        "completed_steps": completed,
    }


def create_tally_record(path: Path, file_name: str, customer: str, aircraft_reg: str,
                        aircraft_type: str, check_type: str, uploader: str,
                        description: str = "", tally_format: str | None = None) -> str:
    adapter = adapter_for_upload(customer, tally_format)
    tally_id = f"TS-{datetime.now():%Y%m%d}-{uuid.uuid4().hex[:5].upper()}"
    with db() as connection:
        connection.execute(
            """INSERT INTO tally_sheet(id,file_name,stored_path,customer,aircraft_reg,aircraft_type,
            check_type,adapter,uploader,uploaded_at,row_count,status,description)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                tally_id, file_name, portable_stored_path(path), customer, aircraft_reg, aircraft_type,
                check_type, adapter, uploader, now(), 0, "Đã tiếp nhận", description.strip(),
            ),
        )
        ensure_workflow_steps(connection, tally_id)
        set_workflow_step(connection, tally_id, "intake", "completed", {
            "file_name": file_name,
            "description": description.strip(),
            "tally_format": adapter,
            "received_at": now(),
        })
    return tally_id


def inspect_tally_record(tally_id: str) -> dict[str, Any]:
    with db() as connection:
        tally = connection.execute("SELECT * FROM tally_sheet WHERE id=?", (tally_id,)).fetchone()
    if not tally:
        raise HTTPException(404, "Không tìm thấy Tally Sheet")
    path = resolve_stored_path(tally["stored_path"], tally["file_name"])
    if not path.exists():
        raise HTTPException(404, "Không tìm thấy file Tally Sheet gốc")
    try:
        parsed = parse_tally(tally["adapter"], path)
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook.active
        structure = inspect_tally_structure(tally["adapter"], path)
        detected_columns: list[str] = []
        for row in worksheet.iter_rows(
            min_row=1, max_row=min(worksheet.max_row, 60), values_only=True
        ):
            columns = [str(value).strip() for value in row if value not in (None, "")]
            normalized = {normalize(value) for value in columns}
            if tally["adapter"] == "VietJet" and "WORK STEP / EVENT NUMBER" in normalized:
                detected_columns = columns
                break
            if tally["adapter"] == "Lufthansa" and "CRI" in normalized:
                detected_columns = columns
                break
        if structure:
            ms_item = next(
                (block for block in structure["blocks"] if block["is_taskcard_scope"]), None
            )
            if ms_item:
                detected_columns = ms_item["fields"]
        task_keys = [normalize(item.get("task_no")) for item in parsed]
        duplicate_count = len(task_keys) - len(set(task_keys))
        warnings = []
        for label, value in (
            ("Loại tàu bay", tally["aircraft_type"]),
            ("Check Type", tally["check_type"]),
        ):
            if not str(value or "").strip():
                warnings.append(f"Chưa có {label}.")
        if duplicate_count:
            warnings.append(f"Có {duplicate_count} mã task trùng trong Tally Sheet.")
        result = {
            "adapter": tally["adapter"],
            "sheet_names": workbook.sheetnames,
            "active_sheet": worksheet.title,
            "source_rows": worksheet.max_row,
            "source_columns": worksheet.max_column,
            "detected_columns": detected_columns,
            "row_count": len(parsed),
            "duplicate_count": duplicate_count,
            "warnings": warnings,
            "confirmed": False,
            "read_at": now(),
            "structure": structure,
            "sample_rows": [
                {
                    "seq": item.get("seq"),
                    "task_no": item.get("task_no"),
                    "title": item.get("source_title"),
                    "section": item.get("section") or item.get("event_type"),
                }
                for item in parsed[:5]
            ],
        }
    except HTTPException:
        raise
    except Exception as exc:
        with db() as connection:
            connection.execute(
                "UPDATE tally_sheet SET status=?,error_message=? WHERE id=?",
                ("Lỗi đọc file", str(exc), tally_id),
            )
            set_workflow_step(
                connection, tally_id, "inspect", "in_progress",
                {"error": str(exc), "read_at": now()},
            )
        raise HTTPException(400, str(exc)) from exc
    with db() as connection:
        connection.execute(
            "UPDATE tally_sheet SET row_count=?,status=?,error_message=NULL WHERE id=?",
            (len(parsed), "Đã đọc sơ bộ", tally_id),
        )
        set_workflow_step(connection, tally_id, "inspect", "in_progress", result)
    return result


def confirm_tally_inspection(tally_id: str) -> dict[str, Any]:
    with db() as connection:
        row = connection.execute(
            "SELECT * FROM workflow_step WHERE tally_id=? AND step_code='inspect'",
            (tally_id,),
        ).fetchone()
        if not row:
            raise HTTPException(400, "Cần chạy đọc sơ bộ trước khi xác nhận")
        details = decode_details(row["details_json"])
        if not details.get("row_count"):
            raise HTTPException(400, "Kết quả đọc sơ bộ chưa hợp lệ")
        details["confirmed"] = True
        details["confirmed_at"] = now()
        set_workflow_step(connection, tally_id, "inspect", "completed", details)
    return details


def automatic_match_fields(item: dict[str, Any], decision: MatchDecision) -> dict[str, Any]:
    amp = decision.record
    if decision.status == AUTO_MATCH and amp:
        important = 1 if "WATCH ITEM" in normalize(amp["classification"]) else 0
        return {
            "matched": 1, "amp_source": amp["source"], "title": amp["title"],
            "description": amp["description"], "phase": amp["phase"],
            "zone": amp["zone"] or item.get("source_zone"), "area": amp["area"],
            "panel": amp["panel"], "special_code": amp["special_code"],
            "resource_type": amp["resource_type"], "classification": amp["classification"],
            "qualification": amp["qualification"], "planning_mhr": mhr_hours(amp["planning_mhr"]),
            "information_skill": amp["information_skill"],
            "requirement_link": amp["requirement_link"],
            "linked_jobcards": amp["linked_jobcards"], "type_code": amp["type_code"],
            "jobcards_title": amp["jobcards_title"], "important_job": important,
            "important_source": "AMOS Watch Item" if important else None,
            "not_in_database": 0,
        }
    return {
        "matched": 0, "amp_source": None, "title": item.get("source_title"),
        "description": item.get("source_note"), "phase": None,
        "zone": item.get("source_zone"), "area": None, "panel": None,
        "special_code": None, "resource_type": None, "classification": None,
        "qualification": None, "planning_mhr": None, "information_skill": None,
        "requirement_link": None, "linked_jobcards": None, "type_code": None,
        "jobcards_title": None, "important_job": 0, "important_source": None,
        "not_in_database": 1 if decision.status == NOT_FOUND else 0,
    }


# Columns of evaluation_item written by the matching (layer 1 result or not)
LAYER1_RESULT_FIELDS = (
    "amp_source", "amos_reference", "amp_refs", "title", "description", "phase", "zone", "area",
    "panel", "special_code", "resource_type", "classification", "qualification", "planning_mhr",
    "information_skill", "requirement_link", "linked_jobcards", "type_code", "jobcards_title",
    "important_job", "important_source",
)


def layer1_fields(result: dict[str, Any], source_zone: Any = None) -> dict[str, Any]:
    """evaluation_item values for a row matched by the MPD-key flow (mpd_matching.match_rows).

    A row with a review note (several Tally rows share one AMP item, row/item count
    differs, LH/RH unknown, record of another customer) is REVIEW_REQUIRED; its
    values and MHR are still filled in. Otherwise AUTO_MATCH.
    """
    values = {field: result.get(field) for field in LAYER1_RESULT_FIELDS}
    values["zone"] = values.get("zone") or source_zone
    review = result.get("match_note")
    values.update({
        "matched": 1,
        "not_in_database": 0,
        "match_status": REVIEW_REQUIRED if review else AUTO_MATCH,
        "match_method": STEP_METHOD.get(result.get("match_rule"), result.get("match_rule")),
        "match_rule": result.get("match_rule"),
        "match_candidate_count": result.get("match_candidate_count") or 1,
        "match_source_id": result.get("match_source_id"),
        "match_note": combined_note(result.get("info_note"), review),
        "match_layer": 1,
        "match_score": 100.0,
        "suggested_task_no": result.get("suggested_task_no"),
        "suggested_title": result.get("suggested_title"),
        "suggested_mhr": result.get("planning_mhr"),
    })
    return values


def unmatched_fields(item: dict[str, Any], decision: MatchDecision) -> dict[str, Any]:
    """evaluation_item values for a row the MPD-key flow did not match (layers 2/3, not found, n/a)."""
    amp = decision.record
    values = automatic_match_fields(item, decision)
    values.update({
        "amos_reference": None,
        "amp_refs": None,
        "match_rule": None,
        "match_status": decision.status,
        "match_method": decision.method,
        "match_candidate_count": decision.candidate_count,
        "match_source_id": amp.get("id") if amp else None,
        "match_note": decision.note,
        "match_layer": decision.layer,
        "match_score": decision.score,
        "suggested_task_no": (amp.get("taskcard_no") or amp.get("task_no")) if amp else None,
        "suggested_title": amp.get("title") if amp else None,
        "suggested_mhr": mhr_hours(amp.get("planning_mhr")) if amp else None,
    })
    return values


def suggestion_decision(connection: sqlite3.Connection, row: dict[str, Any],
                        customer: str, aircraft_type: str) -> MatchDecision:
    """Layers 2–3 of the fork for a row that layer 1 could not match."""
    return lookup_amp(
        connection, row["task_ref"], customer, aircraft_type,
        tally_title_text(row.get("source_title")), include_exact=False,
    )


def populate_evaluation_items(connection: sqlite3.Connection, evaluation_id: str,
                              tally: sqlite3.Row,
                              parsed: list[dict[str, Any]]) -> tuple[int, dict[str, int]]:
    """Layer 1 = MPD-key flow for every eligible row; layers 2–3 only where layer 1 found nothing."""
    matched = 0
    status_counts = {AUTO_MATCH: 0, REVIEW_REQUIRED: 0, NOT_FOUND: 0, NOT_APPLICABLE: 0}
    rows = [prepare_row(tally["adapter"], dict(item)) for item in parsed]
    eligible = [index for index, row in enumerate(rows) if row["eligible"]]
    layer1 = match_rows(
        connection, [rows[index] for index in eligible], tally["customer"], tally["aircraft_type"]
    )
    results = dict(zip(eligible, layer1))
    for index, item in enumerate(rows):
        result = results.get(index)
        if result:
            values = layer1_fields(result, item.get("source_zone"))
            matched += 1
        else:
            if item["eligible"]:
                decision = suggestion_decision(
                    connection, item, tally["customer"], tally["aircraft_type"]
                )
            else:
                decision = MatchDecision(
                    NOT_APPLICABLE,
                    "EVENT_TYPE_FILTER",
                    None,
                    0,
                    "Event không phải Block Taskcard và Title không có mã task nên không có MPD task reference.",
                )
            values = unmatched_fields(item, decision)
        status_counts[values["match_status"]] += 1
        payload = {
            "evaluation_id": evaluation_id,
            "seq": item.get("seq"),
            "section": item.get("section"),
            "task_no": str(item["task_no"]),
            "source_title": item.get("source_title"),
            "event_type": item.get("event_type"),
            "rii": item.get("rii"),
            "source_note": item.get("source_note"),
            "source_zone": item.get("source_zone"),
            "task_ref": item.get("task_ref"),
            "match_key": item.get("match_key") if item["eligible"] else None,
            **values,
            "match_logic_version": MATCH_LOGIC_VERSION,
            "updated_at": now(),
        }
        columns = ",".join(payload)
        connection.execute(
            f"INSERT INTO evaluation_item({columns}) VALUES ({','.join('?' * len(payload))})",
            tuple(payload.values()),
        )
    return matched, status_counts


def evaluate_tally(tally_id: str) -> tuple[str, bool]:
    with db() as connection:
        tally = connection.execute("SELECT * FROM tally_sheet WHERE id=?", (tally_id,)).fetchone()
        if not tally:
            raise HTTPException(404, "Không tìm thấy Tally Sheet")
        existing = connection.execute(
            "SELECT id FROM evaluation WHERE tally_id=?", (tally_id,)
        ).fetchone()
        if existing:
            return existing["id"], True
        inspection = connection.execute(
            "SELECT details_json FROM workflow_step WHERE tally_id=? AND step_code='inspect'",
            (tally_id,),
        ).fetchone()
        inspection_details = decode_details(inspection["details_json"]) if inspection else {}
        if not inspection_details.get("row_count"):
            raise HTTPException(400, "Cần chạy đọc sơ bộ trước khi đánh giá")
        if not inspection_details.get("confirmed"):
            raise HTTPException(400, "Cần xác nhận dữ liệu đọc sơ bộ trước khi đánh giá")
    path = resolve_stored_path(tally["stored_path"], tally["file_name"])
    parsed = parse_tally(tally["adapter"], path)
    evaluation_id = f"EV-{datetime.now():%Y%m%d}-{uuid.uuid4().hex[:5].upper()}"
    with db() as connection:
        existing = connection.execute(
            "SELECT id FROM evaluation WHERE tally_id=?", (tally_id,)
        ).fetchone()
        if existing:
            return existing["id"], True
        connection.execute(
            "INSERT INTO evaluation(id,tally_id,revision,created_at,updated_at) VALUES (?,?,?,?,?)",
            (evaluation_id, tally_id, "REV00", now(), now()),
        )
        matched, status_counts = populate_evaluation_items(
            connection, evaluation_id, tally, parsed
        )
        connection.execute(
            "UPDATE tally_sheet SET matched_count=?,row_count=?,status=? WHERE id=?",
            (matched, len(parsed), "Đã đánh giá sơ bộ", tally_id),
        )
        set_workflow_step(connection, tally_id, "evaluate", "completed", {
            "evaluation_id": evaluation_id,
            "total": len(parsed),
            "matched": matched,
            "unmatched": status_counts[NOT_FOUND],
            "review_required": status_counts[REVIEW_REQUIRED],
            "not_applicable": status_counts[NOT_APPLICABLE],
            "completed_at": now(),
        })
        set_workflow_step(connection, tally_id, "review", "in_progress", {
            "checklist": {"detail": False, "important": False, "database": False},
            "scope": "not_applicable",
        })
        set_workflow_step(connection, tally_id, "finalize", "not_started", {})
    return evaluation_id, False


def migrate_legacy_matching_results() -> None:
    """Re-run matching for evaluations made by an older matching version.

    Evaluations with manual work (review changes in change_log, or MiniDB changes
    made from them in amp_change) are never re-run.
    """
    with db() as connection:
        candidates = connection.execute(
            """SELECT e.id evaluation_id,t.*
            FROM evaluation e JOIN tally_sheet t ON t.id=e.tally_id
            WHERE EXISTS (SELECT 1 FROM evaluation_item i
                          WHERE i.evaluation_id=e.id AND COALESCE(i.match_logic_version,0)<?)
            AND NOT EXISTS (SELECT 1 FROM change_log c WHERE c.evaluation_id=e.id)
            AND NOT EXISTS (SELECT 1 FROM amp_change a WHERE a.evaluation_id=e.id)""",
            (MATCH_LOGIC_VERSION,),
        ).fetchall()
    for tally in candidates:
        path = resolve_stored_path(tally["stored_path"], tally["file_name"])
        if not path.exists():
            sample_path = SAMPLES_DIR / tally["file_name"]
            if not sample_path.exists():
                continue
            path = sample_path
        try:
            parsed = parse_tally(tally["adapter"], path)
        except Exception:
            continue
        with db() as connection:
            changed = connection.execute(
                "SELECT 1 FROM change_log WHERE evaluation_id=? LIMIT 1",
                (tally["evaluation_id"],),
            ).fetchone()
            if changed:
                continue
            connection.execute(
                "DELETE FROM evaluation_item WHERE evaluation_id=?", (tally["evaluation_id"],)
            )
            matched, status_counts = populate_evaluation_items(
                connection, tally["evaluation_id"], tally, parsed
            )
            connection.execute(
                "UPDATE tally_sheet SET matched_count=?,row_count=?,status=? WHERE id=?",
                (matched, len(parsed), "Đã đánh giá sơ bộ", tally["id"]),
            )
            connection.execute(
                "UPDATE evaluation SET updated_at=? WHERE id=?",
                (now(), tally["evaluation_id"]),
            )
            set_workflow_step(connection, tally["id"], "evaluate", "completed", {
                "evaluation_id": tally["evaluation_id"],
                "total": len(parsed),
                "matched": matched,
                "unmatched": status_counts[NOT_FOUND],
                "review_required": status_counts[REVIEW_REQUIRED],
                "not_applicable": status_counts[NOT_APPLICABLE],
                "matching_migrated_at": now(),
            })


# Status texts written by the pre-merge program -> status codes
LEGACY_STATUS_TEXT = {"ĐÃ KHỚP": AUTO_MATCH, "CHƯA KHỚP": NOT_FOUND}


def migrate_legacy_status_labels() -> None:
    """Rows kept from the pre-merge program store 'Đã khớp' / 'Chưa khớp'; convert them to codes.

    A matched row with a Match Note becomes REVIEW_REQUIRED (same rule as layer1_fields).
    """
    with db() as connection:
        rows = connection.execute(
            """SELECT id,matched,match_status,match_note,match_rule FROM evaluation_item
            WHERE match_status IS NOT NULL
            AND match_status NOT IN ('AUTO_MATCH','REVIEW_REQUIRED','NOT_FOUND','NOT_APPLICABLE')"""
        ).fetchall()
        for row in rows:
            status = LEGACY_STATUS_TEXT.get(normalize(row["match_status"]))
            if status is None:
                status = AUTO_MATCH if row["matched"] else NOT_FOUND
            if status == AUTO_MATCH and (row["match_note"] or "").strip() \
                    and not (row["match_note"] or "").startswith("Mã task lấy từ Title"):
                status = REVIEW_REQUIRED
            layer = 1 if status in {AUTO_MATCH, REVIEW_REQUIRED} else None
            connection.execute(
                """UPDATE evaluation_item SET match_status=?, match_layer=COALESCE(match_layer,?),
                match_score=COALESCE(match_score,?), match_method=COALESCE(match_method,?) WHERE id=?""",
                (status, layer, 100.0 if layer else None,
                 STEP_METHOD.get(row["match_rule"]) if layer else None, row["id"]),
            )


def delete_tally(tally_id: str) -> dict[str, Any]:
    """Delete a Tally Sheet with its evaluation, rows, review history, workflow and uploaded file.

    Changes made to the MiniDB from this evaluation (amp_change, added/edited records) are kept.
    """
    with db() as connection:
        tally = connection.execute("SELECT * FROM tally_sheet WHERE id=?", (tally_id,)).fetchone()
        if not tally:
            raise HTTPException(404, "Không tìm thấy hồ sơ")
        for evaluation in connection.execute(
            "SELECT id FROM evaluation WHERE tally_id=?", (tally_id,)
        ).fetchall():
            connection.execute("DELETE FROM change_log WHERE evaluation_id=?", (evaluation["id"],))
            connection.execute("DELETE FROM evaluation_item WHERE evaluation_id=?", (evaluation["id"],))
            connection.execute("DELETE FROM evaluation WHERE id=?", (evaluation["id"],))
        connection.execute("DELETE FROM workflow_step WHERE tally_id=?", (tally_id,))
        connection.execute("DELETE FROM tally_sheet WHERE id=?", (tally_id,))
    path = resolve_stored_path(tally["stored_path"], tally["file_name"])
    uploads = UPLOADS_DIR.resolve()
    if path.is_file() and uploads in path.resolve().parents:
        path.unlink(missing_ok=True)
    return {"ok": True, "tally_id": tally_id}


def seed_samples() -> None:
    with db() as connection:
        if connection.execute("SELECT COUNT(*) FROM tally_sheet").fetchone()[0]:
            return
    samples = [
        ("(Sample) TALLY SHEET WP VJC-A699-24MT2026 REV02.xlsx", "VA-VJ", "VJC-A699", "A320", "24MT"),
        ("(Sample) Workpackage-Summary for D-AIXB_H-25-29102025.xlsx", "DLH", "D-AIXB", "A350", "H-CHECK"),
    ]
    for name, customer, registration, aircraft, check_type in samples:
        source = SAMPLES_DIR / name
        if source.exists():
            destination = UPLOADS_DIR / f"sample-{name}"
            shutil.copy2(source, destination)
            tally_id = create_tally_record(
                destination, name, customer, registration, aircraft,
                check_type, "Dữ liệu mẫu",
            )
            inspect_tally_record(tally_id)
            confirm_tally_inspection(tally_id)
            evaluate_tally(tally_id)


def migrate_workflows() -> None:
    with db() as connection:
        tallies = connection.execute("SELECT * FROM tally_sheet").fetchall()
        for tally in tallies:
            ensure_workflow_steps(connection, tally["id"])
            steps = {
                row["step_code"]: row
                for row in connection.execute(
                    "SELECT * FROM workflow_step WHERE tally_id=?", (tally["id"],)
                )
            }
            if steps["intake"]["status"] != "completed":
                set_workflow_step(connection, tally["id"], "intake", "completed", {
                    "file_name": tally["file_name"],
                    "description": tally["description"] or "",
                    "received_at": tally["uploaded_at"],
                    "migrated": True,
                })
            evaluation = connection.execute(
                "SELECT * FROM evaluation WHERE tally_id=?", (tally["id"],)
            ).fetchone()
            if tally["row_count"] and steps["inspect"]["status"] == "not_started":
                set_workflow_step(connection, tally["id"], "inspect", "completed", {
                    "adapter": tally["adapter"],
                    "row_count": tally["row_count"],
                    "warnings": [],
                    "confirmed": True,
                    "migrated": True,
                })
            if evaluation and steps["evaluate"]["status"] != "completed":
                set_workflow_step(connection, tally["id"], "evaluate", "completed", {
                    "evaluation_id": evaluation["id"],
                    "total": tally["row_count"],
                    "matched": tally["matched_count"],
                    "unmatched": max(0, tally["row_count"] - tally["matched_count"]),
                    "migrated": True,
                })
            if evaluation and steps["review"]["status"] == "not_started":
                set_workflow_step(connection, tally["id"], "review", "in_progress", {
                    "checklist": {"detail": False, "important": False, "database": False},
                    "scope": "not_applicable",
                    "migrated": True,
                })
            if evaluation:
                final_status = (
                    "completed"
                    if evaluation["status"] in {"Chờ phê duyệt", "Hoàn tất"}
                    else "not_started"
                )
                if steps["finalize"]["status"] != final_status:
                    set_workflow_step(connection, tally["id"], "finalize", final_status, {
                        "evaluation_status": evaluation["status"],
                        "migrated": True,
                    })
