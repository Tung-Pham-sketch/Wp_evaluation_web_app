from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .utils import column_map, normalize


def parse_vietjet(path: Path) -> list[dict[str, Any]]:
    worksheet = load_workbook(path, read_only=True, data_only=True).active
    rows = worksheet.iter_rows(values_only=True)
    headers = list(next(rows))
    header_map = column_map(headers)
    required = ["WORK STEP / EVENT NUMBER", "TITLE", "EVENT"]
    if any(key not in header_map for key in required):
        raise ValueError(
            "Không nhận diện được cấu trúc VietJet: thiếu Work Step / Event Number, Title hoặc Event."
        )
    output = []
    for row in rows:
        task = row[header_map["WORK STEP / EVENT NUMBER"]]
        if task in (None, ""):
            continue

        def value(name):
            index = header_map.get(name)
            return row[index] if index is not None and index < len(row) else None

        output.append({
            "seq": value("SEQ. NO."),
            "section": "TALLY SHEET",
            "task_no": task,
            "source_title": value("TITLE"),
            "event_type": value("EVENT"),
            "rii": value("RII"),
            "source_note": value("NOTE"),
        })
    return output


def inspect_lufthansa_structure(path: Path) -> dict[str, Any]:
    """Read every major block in a Lufthansa Workpackage Summary."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    matrix = [list(row) for row in worksheet.iter_rows(values_only=True)]
    headings: list[tuple[int, str]] = []
    for row_number, row in enumerate(matrix, 1):
        first = str(row[0]).strip() if row and row[0] not in (None, "") else ""
        if re.match(r"^\d+\.\s", first) or "MAINTENANCE EVENTS" in normalize(first):
            headings.append((row_number, first))

    blocks: list[dict[str, Any]] = []
    for index, (title_row, raw_title) in enumerate(headings):
        block_end = headings[index + 1][0] - 1 if index + 1 < len(headings) else worksheet.max_row
        header_row = None
        headers: list[str] = []
        for candidate_row in range(title_row + 1, block_end + 1):
            values = matrix[candidate_row - 1]
            nonblank = [value for value in values if value not in (None, "")]
            if len(nonblank) >= 2:
                last_column = max(i for i, value in enumerate(values) if value not in (None, ""))
                headers = [
                    str(value).strip()
                    for value in values[:last_column + 1]
                    if value not in (None, "")
                ]
                header_row = candidate_row
                break

        data_start = (header_row + 1) if header_row else None
        record_count = 0
        if data_start and data_start <= block_end:
            width = max(1, len(headers))
            for row_number in range(data_start, block_end + 1):
                row = matrix[row_number - 1]
                if any(value not in (None, "") for value in row[:width]):
                    record_count += 1

        normalized_title = normalize(raw_title)
        if "REMOVED AFTER PRINTOUT" in normalized_title:
            title, category = "Maintenance Events removed after printout", "Nhật ký thay đổi"
        elif "ADDED AFTER PRINTOUT" in normalized_title:
            title, category = "Maintenance Events added after printout", "Nhật ký thay đổi"
        elif "MS ITEMS" in normalized_title:
            title, category = raw_title, "Taskcard"
        elif "CHECKS" in normalized_title:
            title, category = raw_title, "Checks"
        elif "WORKORDERS" in normalized_title:
            title, category = raw_title, "Workorders"
        elif "ROTABLE" in normalized_title:
            title, category = raw_title, "Rotable"
        elif "REPLACEMENT" in normalized_title:
            title, category = raw_title, "Replacement"
        elif "DOCUMENTS" in normalized_title:
            title, category = raw_title, "Documents"
        elif "JOBCARDS" in normalized_title:
            title, category = raw_title, "Jobcards"
        else:
            title, category = raw_title, "Khác"

        data_range = "—"
        if data_start and record_count:
            data_range = f"A{data_start}:{get_column_letter(max(1, len(headers)))}{block_end}"
        blocks.append({
            "order": index + 1,
            "title": title,
            "raw_title": raw_title,
            "category": category,
            "title_row": title_row,
            "header_row": header_row,
            "row_span": f"Hàng {title_row}–{block_end}",
            "data_range": data_range,
            "record_count": record_count,
            "fields": headers,
            "is_taskcard_scope": "MS ITEMS" in normalized_title,
        })

    ms_item = next((block for block in blocks if block["is_taskcard_scope"]), None)
    return {
        "form_type": "Báo cáo nhiều khối trên một sheet",
        "sheet_count": len(workbook.sheetnames),
        "sheet_names": workbook.sheetnames,
        "active_sheet": worksheet.title,
        "dimension": worksheet.calculate_dimension(),
        "block_count": len(blocks),
        "total_record_count": sum(block["record_count"] for block in blocks),
        "ms_item_count": ms_item["record_count"] if ms_item else 0,
        "blocks": blocks,
        "metadata_note": (
            "Số đăng ký, Check và ngày hồ sơ nằm trong tên file; "
            "không có vùng metadata riêng ở đầu sheet."
        ),
    }


def parse_lufthansa(path: Path) -> list[dict[str, Any]]:
    worksheet = load_workbook(path, read_only=True, data_only=True).active
    matrix = [list(row) for row in worksheet.iter_rows(values_only=True)]
    output: list[dict[str, Any]] = []
    section = ""
    headers: dict[str, int] = {}
    for row in matrix:
        first = str(row[0]).strip() if row and row[0] not in (None, "") else ""
        if "MAINTENANCE EVENTS" in normalize(first):
            headers = {}
            continue
        if re.match(r"^\d+\.\s", first):
            section = first
            headers = {}
            continue
        row_map = column_map(row)
        if "CRI" in row_map and ("TITEL" in row_map or "TITLE" in row_map):
            headers = row_map
            continue
        if ("EVENT ID" in row_map or "WORK STEP / EVENT NUMBER" in row_map) and "CRI" not in row_map:
            headers = {}
            continue
        if not headers or "CRI" not in headers:
            continue
        cri_index = headers["CRI"]
        cri = row[cri_index] if cri_index < len(row) else None
        if cri in (None, "", "CRI"):
            continue

        def value(*names):
            for name in names:
                index = headers.get(name)
                if index is not None and index < len(row):
                    return row[index]
            return None

        output.append({
            "seq": len(output) + 1,
            "section": section or "OTHER",
            "task_no": cri,
            "source_title": value("TITEL", "TITLE", "MS DESCRIPTION"),
            "event_type": "MS Item" if "MS ITEMS" in section.upper() else section,
            "rii": None,
            "source_note": value("MS DESCRIPTION", "DESCRIPTION"),
            "source_zone": value("ZONE"),
            "source_ata": value("ATA"),
        })
    if not output:
        raise ValueError("Không nhận diện được các section có cột CRI trong file Lufthansa.")
    return output
