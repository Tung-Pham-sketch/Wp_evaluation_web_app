# =============================================================================
# minidb.py – AMOS POPUP AND MANUAL MiniDB MAINTENANCE
# -----------------------------------------------------------------------------
# Only the MiniDB (wp_demo.db) is changed – the real AMOS is never touched.
# Every change is written to amp_change and can be exported to Excel.
#
# item_amos_detail(item_id)        data of the AMOS popup
#   - selected records (from amp_refs) or the layer 2/3 suggestion
#   - checks Tally vs AMOS: Task No. (MPD key), Title, Zone, Planning MHR,
#     Customer
#   - other candidates from both sources; "same_amp" = item of the chosen AMP
#     normally given to another Tally row (e.g. ENG #2)
#
# "Thêm vào Database"
#   add_form_defaults()   proposed values + target (VNA or CUSTOMER from plan)
#   add_to_database()     checks REQUIRED_FIELDS, refuses a duplicate natural
#                         key, inserts origin='MANUAL', logs ADD, then
#                         rematch_unmatched(): layer 1 again for the unmatched
#                         rows of the same evaluation and MPD key.
#
# "Sửa bản ghi"
#   edit_amp_record()     LOCKED_FIELDS (= natural key) are ignored; other
#                         changed fields are logged as EDIT, edited=1, and the
#                         evaluation row being reviewed is refreshed.
#
# AMP_FORM_FIELDS   (column, label, used for VNA, used for Customer, multiline)
#                   drives the form on screen (/api/amp-fields).
# REQUIRED_FIELDS   VNA: task_no, aircraft_type, amp_name, amp_date, title
#                   CUSTOMER: task_no, aircraft_type, customer, title
# export_amp_changes()  Excel "Thay đổi MiniDB" of every ADD / EDIT.
# =============================================================================

"""MiniDB maintenance and the AMOS popup.

"Thêm vào Database" and "Sửa bản ghi" change only the MiniDB (wp_demo.db);
the real AMOS system is never touched. Every change is written to amp_change
and can be exported to Excel. Changes are kept when a Tally Sheet is deleted.
"""
from __future__ import annotations

import io
import json
import re
import sqlite3
from datetime import datetime
from typing import Any

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .database import db, now
from .matching import normalize
from .mpd_matching import (
    METHOD_STEP,
    NATURAL_KEY_FIELDS,
    SOURCE_LABEL,
    STEP_CUSTOMER_FIRST,
    STEP_CUSTOMER_SAME,
    STEP_LABEL,
    STEP_VNA,
    aircraft_family,
    amp_candidates,
    customer_candidates,
    customer_codes,
    find_by_natural_key,
    iso_date,
    match_key,
    matching_plan,
    match_rows,
    merge_items,
    mhr_value,
    natural_key,
    prepare_row,
    split_tokens,
    tally_title_text,
)
from .reference_data import AMP_COLUMNS, ORIGIN_MANUAL, derived_amp_fields
from .services import layer1_fields


# Form fields: (column, label, used for VNA, used for Customer, multi-line)
AMP_FORM_FIELDS = [
    ("task_no", "Task No. / Jobcard No.", True, True, False),
    ("aircraft_type", "A/C Type", True, True, False),
    ("customer", "Customer", False, True, False),
    ("effectivity", "Effectivity", False, True, False),
    ("amp_name", "AMP Name / Rev No.", True, False, False),
    ("amp_date", "AMP Date", True, False, False),
    ("mp_name", "MP Name / Issue / Rev", True, False, False),
    ("taskcard_no", "Taskcard No. / Rev", True, False, False),
    ("taskcard_ident", "Taskcard Ident.", True, False, False),
    ("task_type", "Type", True, False, False),
    ("title", "Taskcard Title", True, True, False),
    ("description", "Description", True, True, True),
    ("phase", "Phase", True, True, False),
    ("zone", "Zone", True, True, False),
    ("area", "Area", True, True, False),
    ("panel", "Panel", True, True, False),
    ("special_code", "Special Code", True, True, False),
    ("resource_type", "Resource Type", True, True, False),
    ("classification", "Classification", True, True, False),
    ("type_code", "Type Code", True, False, False),
    ("qualification", "Qualification", True, True, False),
    ("planning_mhr", "Planning MHR (h:mm)", True, True, False),
    ("information_skill", "Information Skill", True, True, False),
    ("information_mhr", "Information MHR (h:mm)", True, True, False),
    ("requirement_link", "Requirement Link / References", True, True, True),
    ("linked_jobcards", "Linked Jobcards", True, True, False),
    ("jobcards_title", "Jobcards Title", True, True, False),
]
# Mandatory fields when adding a record
REQUIRED_FIELDS: dict[str, set[str]] = {
    "VNA": {"task_no", "aircraft_type", "amp_name", "amp_date", "title"},
    "CUSTOMER": {"task_no", "aircraft_type", "customer", "title"},
}
# Identity of a record: cannot be edited (other evaluations are linked to it)
LOCKED_FIELDS = set(NATURAL_KEY_FIELDS)

AMOS_RECORD_FIELDS = ("id", "source", "amp_name", "amp_date", "mp_name", "taskcard_no", "taskcard_ident",
                      "task_no", "mpd_key", "task_type", "aircraft_type", "customer", "effectivity",
                      "title", "description", "phase", "zone", "area", "panel", "special_code",
                      "resource_type", "classification", "type_code", "qualification", "planning_mhr",
                      "information_skill", "information_mhr", "requirement_link", "linked_jobcards",
                      "jobcards_title", "origin", "edited")


def field_definitions() -> dict[str, Any]:
    return {
        "fields": [{"name": n, "label": l, "vna": v, "customer": c, "multiline": m}
                   for n, l, v, c, m in AMP_FORM_FIELDS],
        "required": {k: sorted(v) for k, v in REQUIRED_FIELDS.items()},
        "locked": sorted(LOCKED_FIELDS),
    }


def record_dict(row: sqlite3.Row, selected: bool) -> dict[str, Any]:
    data = {field: row[field] for field in AMOS_RECORD_FIELDS}
    data["source_label"] = SOURCE_LABEL.get(row["source"], row["source"])
    data["mhr_hours"] = mhr_value(row["planning_mhr"])
    data["selected"] = selected
    return data


def _item_with_tally(conn: sqlite3.Connection, item_id: int) -> sqlite3.Row:
    row = conn.execute("""SELECT i.*, t.customer AS tally_customer, t.aircraft_type AS tally_aircraft_type,
      t.adapter AS tally_adapter FROM evaluation_item i JOIN evaluation e ON e.id=i.evaluation_id
      JOIN tally_sheet t ON t.id=e.tally_id WHERE i.id=?""", (item_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Không tìm thấy task")
    return row


# ---------------------------------------------------------------------------
# AMOS popup: Tally row vs AMOS record(s), other candidates
# ---------------------------------------------------------------------------
def item_amos_detail(item_id: int) -> dict[str, Any]:
    with db() as conn:
        row = _item_with_tally(conn, item_id)
        suggestion = None
        if not row["matched"]:
            if row["match_layer"] in (2, 3) and row["match_source_id"]:
                found = conn.execute("SELECT * FROM amp_task WHERE id=?", (row["match_source_id"],)).fetchone()
                suggestion = found
            if not suggestion:
                raise HTTPException(400, "Dòng này chưa khớp với AMOS")
        if row["matched"] and not row["amp_refs"]:
            raise HTTPException(400, "Dòng này được đánh giá bằng phiên bản cũ. Hãy chạy lại đánh giá để xem bản ghi AMOS.")
        selected, missing = [], []
        if suggestion is None:
            for ref in json.loads(row["amp_refs"]):
                found = find_by_natural_key(conn, ref)
                (selected.append(found) if found else missing.append(ref))
            source = json.loads(row["amp_refs"])[0].get("source")
        else:
            selected = [suggestion]
            source = suggestion["source"]
        task_ref = row["task_ref"] or row["task_no"]
        rule = row["match_rule"] or METHOD_STEP.get(row["match_method"]) \
            or (STEP_VNA if source == "VNA" else STEP_CUSTOMER_SAME)
        family = aircraft_family(row["tally_aircraft_type"])
        codes = customer_codes(row["tally_customer"])
        # Alternatives from both sources, the source that was used first
        vna = amp_candidates(conn, task_ref, family)
        cust = customer_candidates(conn, task_ref, None, None)
        exact = conn.execute("SELECT * FROM amp_task WHERE source='CUSTOMER' AND task_key=?",
                             (normalize(task_ref),)).fetchall()
        cust = list({c["id"]: c for c in [*cust, *exact]}.values())
        cust.sort(key=lambda c: (normalize(c["customer"]) not in codes, str(c["task_no"]), c["id"]))
        candidates = [*vna, *cust] if source == "VNA" else [*cust, *vna]
    selected_ids = {r["id"] for r in selected}
    chosen_amps = {r["amp_name"] for r in selected if r["source"] == "VNA"}
    others = []
    for c in candidates:
        if c["id"] in selected_ids:
            continue
        other = record_dict(c, False)
        # Same group as the chosen record: an item normally given to another Tally row (e.g. ENG #1 / ENG #2)
        if rule == STEP_CUSTOMER_FIRST:
            other["same_amp"] = c["source"] == "CUSTOMER" and normalize(c["customer"]) in codes \
                and (not family or c["aircraft_family"] == family)
        else:
            other["same_amp"] = c["source"] == "VNA" and c["amp_name"] in chosen_amps
        others.append(other)

    # Tally vs AMOS comparison
    amos_tasks = sorted({str(r["task_no"]) for r in selected})
    tally_key = match_key(task_ref)
    task_ok = all(match_key(t) == tally_key for t in amos_tasks) if amos_tasks else None
    tally_title = tally_title_text(row["source_title"])
    amos_titles = sorted({str(r["title"] or "") for r in selected if r["title"]})
    title_ok = None
    if tally_title and amos_titles:
        def words(text: str) -> set[str]:
            return set(re.findall(r"[A-Z0-9]+", normalize(text)))
        # Same title if one contains the other, or most words of the shorter title are shared
        title_ok = any(normalize(a) in normalize(tally_title) or normalize(tally_title) in normalize(a)
                       or len(words(a) & words(tally_title)) >= 0.6 * min(len(words(a)), len(words(tally_title)))
                       for a in amos_titles)
    amos_zones = sorted({z for r in selected for z in split_tokens(r["zone"])})
    tally_zones = sorted(set(split_tokens(row["source_zone"])))
    zone_ok = (set(tally_zones) == set(amos_zones)) if tally_zones else None
    amos_mhr = round(sum(mhr_value(r["planning_mhr"]) for r in selected), 2)
    tally_mhr = row["planning_mhr"] if row["matched"] else None
    mhr_ok = abs((tally_mhr or 0) - amos_mhr) < 0.01 if row["matched"] else None
    checks = [
        {"field": "Task No.", "tally": task_ref, "amos": ", ".join(amos_tasks), "ok": task_ok,
         "hint": f"MPD key: {tally_key}"},
        {"field": "Title", "tally": tally_title or None, "amos": " | ".join(amos_titles) or None, "ok": title_ok},
        {"field": "Zone", "tally": ", ".join(tally_zones) or None, "amos": ", ".join(amos_zones) or None,
         "ok": zone_ok, "hint": None if tally_zones else "Tally Sheet không có cột Zone"},
        {"field": "Planning MHR (giờ)", "tally": tally_mhr, "amos": amos_mhr, "ok": mhr_ok,
         "hint": ("Gợi ý lớp 2/3 – MHR chưa được gán" if suggestion is not None
                  else f"Tổng {len(selected)} item" if len(selected) > 1 else None)},
    ]
    if source == "CUSTOMER" and selected:
        same = normalize(selected[0]["customer"]) in codes
        checks.append({"field": "Customer", "tally": row["tally_customer"], "amos": selected[0]["customer"],
                       "ok": same})
    return {
        "item": {key: row[key] for key in ("id", "task_no", "task_ref", "source_title", "event_type",
                                           "matched", "match_status", "match_note", "amos_reference",
                                           "match_layer", "match_score", "planning_mhr", "reviewed_mhr")},
        "checks": checks,
        "selected": [record_dict(r, True) for r in selected],
        "is_suggestion": suggestion is not None,
        "missing": missing,
        "others": others,
        "rule": STEP_LABEL.get(rule, row["match_note"] or rule),
        "plan": [STEP_LABEL[s] for s in matching_plan(row["tally_customer"], row["tally_aircraft_type"])],
    }


# ---------------------------------------------------------------------------
# Add / edit MiniDB records
# ---------------------------------------------------------------------------
def clean_form_values(raw: dict[str, Any], target: str) -> dict[str, Any]:
    allowed = {f[0] for f in AMP_FORM_FIELDS if (f[2] if target == "VNA" else f[3])}
    values = {}
    for key, value in raw.items():
        if key in allowed:
            text = str(value).strip() if value is not None else ""
            values[key] = text or None
    if values.get("amp_date"):
        values["amp_date"] = iso_date(values["amp_date"]) or values["amp_date"]
    if values.get("customer"):
        values["customer"] = normalize(values["customer"])
    return values


def log_amp_change(conn: sqlite3.Connection, action: str, record: sqlite3.Row, field: str,
                   old: Any, new: Any, reason: str, changed_by: str, item: sqlite3.Row | None) -> None:
    conn.execute("""INSERT INTO amp_change(action,amp_id,source,record_key,task_no,field_name,old_value,new_value,
                    reason,changed_by,changed_at,evaluation_id,tally_task_no) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                 (action, record["id"], record["source"], json.dumps(natural_key(record), ensure_ascii=False),
                  record["task_no"], field, None if old is None else str(old), None if new is None else str(new),
                  reason, changed_by, now(), item["evaluation_id"] if item else None,
                  item["task_no"] if item else None))


def apply_layer1(conn: sqlite3.Connection, item_id: int, values: dict[str, Any]) -> None:
    values = {**values, "updated_at": now()}
    sets = ",".join(f"{k}=?" for k in values)
    conn.execute(f"UPDATE evaluation_item SET {sets} WHERE id=?", (*values.values(), item_id))


def refresh_matched_count(conn: sqlite3.Connection, evaluation_id: str) -> None:
    conn.execute("UPDATE tally_sheet SET matched_count=(SELECT COUNT(*) FROM evaluation_item i WHERE "
                 "i.evaluation_id=? AND i.matched=1) WHERE id=(SELECT tally_id FROM evaluation WHERE id=?)",
                 (evaluation_id, evaluation_id))


def rematch_unmatched(conn: sqlite3.Connection, evaluation_id: str, key: str | None) -> int:
    """Run layer 1 again for the unmatched rows of one evaluation that share an MPD key."""
    if not key:
        return 0
    tally = conn.execute("""SELECT t.customer, t.aircraft_type, t.adapter FROM evaluation e
        JOIN tally_sheet t ON t.id=e.tally_id WHERE e.id=?""", (evaluation_id,)).fetchone()
    items = conn.execute("""SELECT * FROM evaluation_item WHERE evaluation_id=? AND matched=0
        AND match_key=? ORDER BY id""", (evaluation_id, key)).fetchall()
    rows = [prepare_row(tally["adapter"], {"task_no": it["task_no"], "source_title": it["source_title"],
                                           "event_type": it["event_type"]}) for it in items]
    count = 0
    for it, result in zip(items, match_rows(conn, rows, tally["customer"], tally["aircraft_type"])):
        if not result:
            continue
        count += 1
        apply_layer1(conn, it["id"], {**layer1_fields(result, it["source_zone"]), "added_to_db": 1})
    refresh_matched_count(conn, evaluation_id)
    return count


def add_form_defaults(item_id: int) -> dict[str, Any]:
    """Values proposed in the 'Thêm vào Database' form, taken from the Tally row."""
    with db() as conn:
        row = _item_with_tally(conn, item_id)
    plan = matching_plan(row["tally_customer"], row["tally_aircraft_type"])
    return {
        "target": "CUSTOMER" if plan[0] == STEP_CUSTOMER_FIRST else "VNA",
        "values": {"task_no": row["task_ref"] or row["task_no"], "title": tally_title_text(row["source_title"]),
                   "aircraft_type": row["tally_aircraft_type"], "customer": normalize(row["tally_customer"]),
                   "zone": row["source_zone"], "description": row["source_note"]},
    }


def add_to_database(item_id: int, target: str, raw_values: dict[str, Any], reason: str,
                    changed_by: str) -> dict[str, Any]:
    target = target.upper()
    if target not in REQUIRED_FIELDS:
        raise HTTPException(400, "Chọn nơi thêm: VNA (ALL AMP Task) hoặc Customer AMP")
    if not reason.strip():
        raise HTTPException(400, "Bắt buộc nhập lý do thay đổi")
    values = clean_form_values(raw_values, target)
    missing = [label for name, label, *_ in AMP_FORM_FIELDS if name in REQUIRED_FIELDS[target] and not values.get(name)]
    if missing:
        raise HTTPException(400, "Bắt buộc nhập: " + ", ".join(missing))
    with db() as conn:
        item = conn.execute("SELECT * FROM evaluation_item WHERE id=?", (item_id,)).fetchone()
        if not item:
            raise HTTPException(404, "Không tìm thấy task")
        if item["matched"]:
            raise HTTPException(400, "Dòng này đã khớp. Dùng chức năng sửa bản ghi AMOS.")
        record = {c: None for c in AMP_COLUMNS}
        record.update(values)
        record.update({"source": target, "origin": ORIGIN_MANUAL})
        record.update(derived_amp_fields(record))
        if conn.execute("SELECT 1 FROM amp_task WHERE " +
                        " AND ".join(f"{f} IS ?" for f in NATURAL_KEY_FIELDS),
                        tuple(record.get(f) for f in NATURAL_KEY_FIELDS)).fetchone():
            raise HTTPException(400, "Bản ghi với cùng thông tin nhận diện đã có trong MiniDB")
        columns = list(AMP_COLUMNS) + ["edited", "updated_at", "updated_by"]
        row_values = [record.get(c) for c in AMP_COLUMNS] + [0, now(), changed_by]
        cursor = conn.execute(f"INSERT INTO amp_task({','.join(columns)}) VALUES ({','.join('?' * len(columns))})",
                              row_values)
        new = conn.execute("SELECT * FROM amp_task WHERE id=?", (cursor.lastrowid,)).fetchone()
        for name, *_ in AMP_FORM_FIELDS:
            if values.get(name) is not None:
                log_amp_change(conn, "ADD", new, name, None, values[name], reason, changed_by, item)
        matched = rematch_unmatched(conn, item["evaluation_id"], item["match_key"])
        conn.execute("UPDATE evaluation_item SET added_to_db=1, updated_at=? WHERE id=?", (now(), item_id))
        conn.execute("UPDATE evaluation SET updated_at=? WHERE id=?", (now(), item["evaluation_id"]))
    return {"ok": True, "amp_id": new["id"], "rematched": matched}


def edit_amp_record(amp_id: int, raw_values: dict[str, Any], reason: str, item_id: int | None,
                    changed_by: str) -> dict[str, Any]:
    """Edit a MiniDB record; new values override the old ones. Identity fields are locked."""
    if not reason.strip():
        raise HTTPException(400, "Bắt buộc nhập lý do thay đổi")
    with db() as conn:
        record = conn.execute("SELECT * FROM amp_task WHERE id=?", (amp_id,)).fetchone()
        if not record:
            raise HTTPException(404, "Không tìm thấy bản ghi")
        values = clean_form_values(raw_values, record["source"])
        changes = {k: v for k, v in values.items()
                   if k not in LOCKED_FIELDS and (record[k] or None) != v}
        if not changes:
            return {"ok": True, "changed": 0}
        item = conn.execute("SELECT * FROM evaluation_item WHERE id=?", (item_id,)).fetchone() \
            if item_id else None
        for field, value in changes.items():
            log_amp_change(conn, "EDIT", record, field, record[field], value, reason, changed_by, item)
        changed_count = len(changes)
        merged = {**dict(record), **changes}
        changes["side"] = derived_amp_fields(merged)["side"]
        changes.update({"edited": 1, "updated_at": now(), "updated_by": changed_by})
        sets = ",".join(f"{k}=?" for k in changes)
        conn.execute(f"UPDATE amp_task SET {sets} WHERE id=?", (*changes.values(), amp_id))
        if item and item["amp_refs"]:   # refresh the evaluation row that is being reviewed
            records = [r for r in (find_by_natural_key(conn, ref) for ref in json.loads(item["amp_refs"])) if r]
            if records:
                merged_values = merge_items(records)
                keep = {k: merged_values[k] for k in (
                    "amp_source", "amos_reference", "amp_refs", "title", "description", "phase", "zone",
                    "area", "panel", "special_code", "resource_type", "classification", "qualification",
                    "planning_mhr", "information_skill", "requirement_link", "linked_jobcards", "type_code",
                    "jobcards_title", "important_job", "important_source")}
                keep["zone"] = keep["zone"] or item["source_zone"]
                keep["suggested_mhr"] = keep["planning_mhr"]
                keep["suggested_title"] = merged_values["suggested_title"]
                apply_layer1(conn, item["id"], keep)
                conn.execute("UPDATE evaluation SET updated_at=? WHERE id=?", (now(), item["evaluation_id"]))
    return {"ok": True, "changed": changed_count}


def get_amp_record(amp_id: int) -> dict[str, Any]:
    with db() as conn:
        record = conn.execute("SELECT * FROM amp_task WHERE id=?", (amp_id,)).fetchone()
    if not record:
        raise HTTPException(404, "Không tìm thấy bản ghi")
    return dict(record)


def _style_sheet(worksheet) -> None:
    for cell in worksheet[1]:
        cell.fill = PatternFill("solid", fgColor="003B71")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column in range(1, worksheet.max_column + 1):
        width = min(42, max(12, max(len(str(worksheet.cell(row, column).value or ""))
                                    for row in range(1, min(worksheet.max_row, 80) + 1)) + 2))
        worksheet.column_dimensions[get_column_letter(column)].width = width
    worksheet.sheet_properties.tabColor = "C89B3C"


def export_amp_changes() -> tuple[io.BytesIO, str]:
    """All manual changes to the MiniDB, for the AMOS data owners or to re-apply later."""
    with db() as conn:
        rows = conn.execute("SELECT * FROM amp_change ORDER BY id").fetchall()
    labels = {n: l for n, l, *_ in AMP_FORM_FIELDS}
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Thay đổi MiniDB"
    worksheet.append(["Thời gian", "Hành động", "Nguồn", "Task No.", "Trường", "Giá trị cũ", "Giá trị mới",
                      "Lý do", "Người sửa", "Mã đánh giá", "Task trong Tally", "Khóa bản ghi"])
    for r in rows:
        worksheet.append([r["changed_at"], "Thêm mới" if r["action"] == "ADD" else "Sửa",
                          SOURCE_LABEL.get(r["source"], r["source"]), r["task_no"],
                          labels.get(r["field_name"], r["field_name"]), r["old_value"], r["new_value"],
                          r["reason"], r["changed_by"], r["evaluation_id"], r["tally_task_no"], r["record_key"]])
    _style_sheet(worksheet)
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream, f"MiniDB_changes_{datetime.now():%Y%m%d}.xlsx"
