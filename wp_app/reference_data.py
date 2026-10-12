from __future__ import annotations

import json
from typing import Any

from openpyxl import load_workbook

from .config import AMOS_DIR
from .database import db
from .matching import AUTO_MATCH, aircraft_from_amp_name, canonical_aircraft, taskcard_core
from .mpd_matching import aircraft_family, base_key, find_by_natural_key, iso_date, mpd_key, side_of
from .utils import column_map, normalize


ORIGIN_AMOS = "AMOS"        # imported from the AMOS Excel exports
ORIGIN_MANUAL = "MANUAL"    # added by a user with "Thêm vào Database"

# Every amp_task column filled by the import (fork columns + MPD matching columns)
AMP_COLUMNS = (
    "source", "aircraft_type", "aircraft_family", "aircraft_model", "customer", "effectivity",
    "amp_name", "amp_date", "mp_name", "taskcard_no", "taskcard_key", "taskcard_core_key",
    "taskcard_ident", "task_type", "task_no", "task_key", "task_base", "mpd_key", "side",
    "title", "description", "phase", "zone", "area", "panel", "special_code",
    "resource_type", "classification", "type_code", "qualification", "planning_mhr",
    "information_skill", "information_mhr", "requirement_link", "linked_jobcards", "jobcards_title",
    "origin",
)


def derived_amp_fields(values: dict[str, Any]) -> dict[str, Any]:
    """Columns the program derives itself from task number, AMP Name, A/C Type, Ident./Area."""
    task_no = values.get("task_no")
    taskcard = values.get("taskcard_no") or task_no
    if values.get("source") == "VNA":
        model = aircraft_from_amp_name(values.get("amp_name"), values.get("aircraft_type"))
    else:
        model = canonical_aircraft(values.get("aircraft_type"))
    return {
        "aircraft_family": aircraft_family(values.get("aircraft_type")),
        "aircraft_model": model,
        "task_key": normalize(task_no),
        "task_base": base_key(task_no),
        "mpd_key": mpd_key(task_no),
        "taskcard_key": normalize(taskcard),
        "taskcard_core_key": taskcard_core(taskcard),
        "side": side_of(values.get("taskcard_ident"), values.get("area")),
    }


def import_amos() -> None:
    """Load ALL AMP Task + ALL Customer AMP into the MiniDB.

    Runs once; runs again only when the MiniDB was built by an older version
    (missing columns). Records added by hand (origin MANUAL) are kept, and
    manual edits are re-applied from amp_change after the rebuild.
    """
    with db() as connection:
        total = connection.execute(
            "SELECT COUNT(*) FROM amp_task WHERE COALESCE(origin,'')<>?", (ORIGIN_MANUAL,)
        ).fetchone()[0]
        outdated = connection.execute(
            """SELECT COUNT(*) FROM amp_task WHERE COALESCE(origin,'')<>? AND (
                 origin IS NULL OR taskcard_core_key IS NULL OR aircraft_model IS NULL
                 OR (source='VNA' AND amp_name IS NULL))""",
            (ORIGIN_MANUAL,),
        ).fetchone()[0]
        if total and not outdated:
            return
        rebuilt = bool(total)
        if total:
            connection.execute("DELETE FROM amp_task WHERE COALESCE(origin,'')<>?", (ORIGIN_MANUAL,))
    sources = [
        (AMOS_DIR / "ALL AMP Task.xlsx", "VNA", "AMP"),
        (AMOS_DIR / "ALL Customer AMP.xlsx", "CUSTOMER", None),
    ]
    records: list[tuple] = []
    for path, source, sheet in sources:
        if not path.exists():
            continue
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook[sheet] if sheet and sheet in workbook.sheetnames else workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = list(next(rows))
        header_map = column_map(headers)

        def value(row, *names):
            for name in names:
                index = header_map.get(normalize(name))
                if index is not None and index < len(row):
                    return row[index]
            return None

        for row in rows:
            task_no = value(row, "Task No./Rev No.", "Jobcard No.", "Taskcard No./Rev")
            if not task_no:
                continue
            values = {
                "source": source,
                "aircraft_type": value(row, "A/C Type"),
                "customer": value(row, "Customer"),
                "effectivity": value(row, "Effectivity"),
                "amp_name": value(row, "AMP Name/ Rev No.", "AMP Name/Rev No."),
                "amp_date": iso_date(value(row, "AMP Date")),
                "mp_name": value(row, "MP Name/Issue No./Rev No."),
                "taskcard_no": value(row, "Taskcard No./Rev"),
                "taskcard_ident": value(row, "Taskcard Ident."),
                "task_type": value(row, "Type"),
                "task_no": str(task_no).strip(),
                "title": value(row, "Taskcard Title", "Title"),
                "description": value(row, "Description"),
                "phase": value(row, "Phase"),
                "zone": value(row, "Zone"),
                "area": value(row, "Area"),
                "panel": value(row, "Panel"),
                "special_code": value(row, "Special Code"),
                "resource_type": value(row, "Resource Type"),
                "classification": value(row, "Classification"),
                "type_code": value(row, "Type Code"),
                "qualification": value(row, "Qualification"),
                "planning_mhr": value(row, "Planning MHR"),
                "information_skill": value(row, "Information Skill"),
                "information_mhr": value(row, "Information MHR"),
                "requirement_link": value(row, "Requirement Link", "References"),
                "linked_jobcards": value(row, "Linked Jobcards"),
                "jobcards_title": value(row, "Jobcards Title"),
                "origin": ORIGIN_AMOS,
            }
            if values["taskcard_no"] is not None:
                values["taskcard_no"] = str(values["taskcard_no"]).strip()
            values.update(derived_amp_fields(values))
            records.append(tuple(values.get(column) for column in AMP_COLUMNS))
    with db() as connection:
        connection.executemany(
            f"INSERT INTO amp_task({','.join(AMP_COLUMNS)}) VALUES ({','.join('?' * len(AMP_COLUMNS))})",
            records,
        )
        if rebuilt:
            reapply_manual_edits(connection)


def reapply_manual_edits(connection) -> None:
    """After a rebuild, apply the logged "Sửa bản ghi" changes again (matched by natural key)."""
    rows = connection.execute(
        "SELECT * FROM amp_change WHERE action='EDIT' ORDER BY id"
    ).fetchall()
    for change in rows:
        try:
            key = json.loads(change["record_key"] or "{}")
        except json.JSONDecodeError:
            continue
        record = find_by_natural_key(connection, key)
        if not record or change["field_name"] not in record.keys():
            continue
        connection.execute(
            f"UPDATE amp_task SET {change['field_name']}=?, edited=1, updated_at=?, updated_by=? WHERE id=?",
            (change["new_value"], change["changed_at"], change["changed_by"], record["id"]),
        )
        updated = dict(connection.execute("SELECT * FROM amp_task WHERE id=?", (record["id"],)).fetchone())
        derived = derived_amp_fields(updated)
        connection.execute("UPDATE amp_task SET side=? WHERE id=?", (derived["side"], record["id"]))


def import_jobcards() -> None:
    with db() as connection:
        if connection.execute("SELECT COUNT(*) FROM jobcard_master").fetchone()[0]:
            return
    path = AMOS_DIR / "All Jobcard.xlsx"
    if not path.exists():
        return
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = workbook.active.iter_rows(values_only=True)
    header_map = column_map(list(next(rows)))

    def value(row, name):
        index = header_map.get(normalize(name))
        return row[index] if index is not None and index < len(row) else None

    records: list[tuple[Any, ...]] = []
    for row in rows:
        jobcard_no = value(row, "Jobcard No.")
        if not jobcard_no:
            continue
        meta_key = value(row, "Meta Key")
        meta_value = value(row, "Meta Value")
        aircraft_type = meta_value if normalize(meta_key) == "AIRCRAFT TYPE" else None
        records.append((
            meta_key, meta_value, aircraft_type, str(jobcard_no).strip(), normalize(jobcard_no),
            value(row, "Title"), value(row, "Description"), value(row, "Phase"),
            value(row, "Zone"), value(row, "Area"), value(row, "Panel"),
            value(row, "Special Code"), value(row, "Resource Type"), value(row, "Classification"),
            value(row, "Qualification"), value(row, "Planning MHR"), value(row, "Information Skill"),
            value(row, "Information MHR"), value(row, "References"),
        ))
    with db() as connection:
        connection.executemany(
            """INSERT INTO jobcard_master(meta_key,meta_value,aircraft_type,jobcard_no,
            jobcard_key,title,description,phase,zone,area,panel,special_code,resource_type,
            classification,qualification,planning_mhr,information_skill,information_mhr,
            requirement_link) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            records,
        )


def backfill_output_columns() -> None:
    """Fork legacy fix: fill Type Code / Jobcards Title of rows matched by an old version.

    Rows matched by the MPD-key flow (amp_refs set) already carry these values and are skipped.
    """
    from .services import lookup_amp

    with db() as connection:
        pending = connection.execute("""SELECT i.id,i.task_no,t.customer,t.aircraft_type
          FROM evaluation_item i JOIN evaluation e ON e.id=i.evaluation_id
          JOIN tally_sheet t ON t.id=e.tally_id
          WHERE i.matched=1 AND i.amp_refs IS NULL
          AND i.type_code IS NULL AND i.jobcards_title IS NULL""").fetchall()
        for item in pending:
            decision = lookup_amp(connection, item["task_no"], item["customer"], item["aircraft_type"])
            amp = decision.record
            if decision.status == AUTO_MATCH and amp:
                connection.execute(
                    "UPDATE evaluation_item SET type_code=?,jobcards_title=? WHERE id=?",
                    (amp["type_code"], amp["jobcards_title"], item["id"]),
                )
