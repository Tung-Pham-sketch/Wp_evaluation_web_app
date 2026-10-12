from __future__ import annotations

import io
from typing import Any

from fastapi import APIRouter, HTTPException
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel

from ..database import db, now
from ..services import resolve_stored_path, set_workflow_step, workflow_progress


router = APIRouter()

# Merged: how the status codes stored in the database are shown to users
STATUS_LABELS = {
    "AUTO_MATCH": "Đã khớp",
    "REVIEW_REQUIRED": "Cần review",
    "NOT_FOUND": "Chưa khớp",
    "NOT_APPLICABLE": "Không áp dụng",
}


@router.get("/api/evaluations/{evaluation_id}")
def evaluation_detail(evaluation_id: str, q: str = "", status: str = "all", page: int = 1,
                      page_size: int = 50, sort_by: str = "seq", sort_dir: str = "asc"):
    page = max(1, page)
    page_size = min(max(10, page_size), 200)
    where = ["i.evaluation_id=?"]
    params: list[Any] = [evaluation_id]
    if q:
        where.append("""(i.task_no LIKE ? OR i.source_title LIKE ? OR i.title LIKE ?
          OR i.description LIKE ? OR i.suggested_task_no LIKE ? OR i.suggested_title LIKE ?
          OR i.amos_reference LIKE ? OR i.task_ref LIKE ?)""")
        params.extend([f"%{q}%"] * 8)
    if status == "matched":
        where.append("i.matched=1")
    if status == "auto_match":
        where.append("i.match_status='AUTO_MATCH'")
    if status == "unmatched":
        where.append("i.matched=0")
    if status == "not_found":
        where.append("i.match_status='NOT_FOUND'")
    if status in {"review_required", "review"}:
        where.append("i.match_status='REVIEW_REQUIRED'")
    if status == "layer_1_review":
        where.append("i.match_status='REVIEW_REQUIRED' AND i.match_layer=1")
    if status == "layer_2":
        where.append("i.match_layer=2")
    if status == "layer_3":
        where.append("i.match_layer=3")
    if status == "not_applicable":
        where.append("i.match_status='NOT_APPLICABLE'")
    if status == "important":
        where.append("i.important_job=1")
    if status == "added_to_db":
        where.append("COALESCE(i.added_to_db,0)=1")
    predicate = " AND ".join(where)
    with db() as connection:
        header = connection.execute("""SELECT e.*,t.file_name,t.description,t.customer,
          t.aircraft_reg,t.aircraft_type,t.check_type,t.adapter,t.version,t.uploaded_at,t.uploader
          FROM evaluation e JOIN tally_sheet t ON t.id=e.tally_id WHERE e.id=?""",
            (evaluation_id,),
        ).fetchone()
        if not header:
            raise HTTPException(404, "Không tìm thấy lần đánh giá")
        total = connection.execute(
            f"SELECT COUNT(*) FROM evaluation_item i WHERE {predicate}", params
        ).fetchone()[0]
        sortable = {
            "task_no", "title", "description", "phase", "zone", "area", "panel",
            "special_code", "resource_type", "classification", "type_code",
            "qualification", "planning_mhr", "information_skill", "requirement_link",
            "linked_jobcards", "jobcards_title", "important_job", "not_in_database",
            "match_status", "match_method", "match_candidate_count", "match_layer",
            "match_score", "suggested_task_no", "suggested_title", "suggested_mhr", "seq",
            "source_title", "match_note", "amos_reference", "task_ref",
        }
        order_column = sort_by if sort_by in sortable else "seq"
        direction = "DESC" if sort_dir.lower() == "desc" else "ASC"
        items = [
            dict(row)
            for row in connection.execute(
                f"""SELECT i.* FROM evaluation_item i WHERE {predicate}
                ORDER BY i.{order_column} {direction},i.id ASC LIMIT ? OFFSET ?""",
                (*params, page_size, (page - 1) * page_size),
            )
        ]
        summary = dict(connection.execute("""SELECT COUNT(*) total,
          COALESCE(SUM(matched),0) matched,
          COALESCE(SUM(not_in_database),0) unmatched,
          COALESCE(SUM(important_job),0) important,
          SUM(CASE WHEN match_status='AUTO_MATCH' THEN 1 ELSE 0 END) auto_match,
          SUM(CASE WHEN match_status='REVIEW_REQUIRED' THEN 1 ELSE 0 END) review_required,
          SUM(CASE WHEN match_status='REVIEW_REQUIRED' AND match_layer=1 THEN 1 ELSE 0 END) layer_1_review,
          SUM(CASE WHEN match_status='NOT_APPLICABLE' THEN 1 ELSE 0 END) not_applicable,
          SUM(CASE WHEN match_layer=1 THEN 1 ELSE 0 END) layer_1,
          SUM(CASE WHEN match_layer=2 THEN 1 ELSE 0 END) layer_2,
          SUM(CASE WHEN match_layer=3 THEN 1 ELSE 0 END) layer_3,
          ROUND(COALESCE(SUM(COALESCE(reviewed_mhr,planning_mhr)),0),2) total_mhr
          FROM evaluation_item WHERE evaluation_id=?""", (evaluation_id,)).fetchone())
        special = [
            dict(row)
            for row in connection.execute("""SELECT COALESCE(NULLIF(special_code,''),'Chưa phân loại') special_code,
              COUNT(*) task_count,ROUND(SUM(COALESCE(reviewed_mhr,planning_mhr)),2) mhr
              FROM evaluation_item WHERE evaluation_id=? GROUP BY 1 ORDER BY mhr DESC LIMIT 12""",
                (evaluation_id,),
            )
        ]
    return {
        "header": dict(header),
        "summary": summary,
        "items": items,
        "special_codes": special,
        "filtered_total": total,
        "page": page,
        "page_size": page_size,
    }


class ItemUpdate(BaseModel):
    reviewed_mhr: float | None = None
    important_job: bool | None = None
    review_note: str | None = None
    reason: str
    changed_by: str = "Demo User"


@router.patch("/api/items/{item_id}")
def update_item(item_id: int, payload: ItemUpdate):
    if not payload.reason.strip():
        raise HTTPException(400, "Bắt buộc nhập lý do thay đổi")
    with db() as connection:
        row = connection.execute("SELECT * FROM evaluation_item WHERE id=?", (item_id,)).fetchone()
        if not row:
            raise HTTPException(404, "Không tìm thấy task")
        changes = {
            key: value
            for key, value in {
                "reviewed_mhr": payload.reviewed_mhr,
                "important_job": int(payload.important_job) if payload.important_job is not None else None,
                "review_note": payload.review_note,
            }.items()
            if value is not None
        }
        for field, value in changes.items():
            if str(row[field] or "") != str(value):
                connection.execute(
                    """INSERT INTO change_log(evaluation_id,item_id,field_name,old_value,
                    new_value,reason,changed_by,changed_at) VALUES (?,?,?,?,?,?,?,?)""",
                    (
                        row["evaluation_id"], item_id, field, row[field], value,
                        payload.reason, payload.changed_by, now(),
                    ),
                )
                connection.execute(
                    f"UPDATE evaluation_item SET {field}=?,updated_at=? WHERE id=?",
                    (value, now(), item_id),
                )
        connection.execute(
            "UPDATE evaluation SET updated_at=? WHERE id=?", (now(), row["evaluation_id"])
        )
    return {"ok": True}


class EvaluationUpdate(BaseModel):
    tat_review: str | None = None
    assumptions: str | None = None
    notes: str | None = None
    prepared_by: str | None = None
    approved_by: str | None = None
    status: str | None = None


@router.patch("/api/evaluations/{evaluation_id}")
def update_evaluation(evaluation_id: str, payload: EvaluationUpdate):
    values = payload.model_dump(exclude_none=True)
    if payload.status and payload.status not in {"Đang review", "Chờ phê duyệt", "Hoàn tất"}:
        raise HTTPException(400, "Trạng thái hồ sơ không hợp lệ")
    if values:
        fields = ",".join(f"{key}=?" for key in values)
        with db() as connection:
            evaluation = connection.execute(
                "SELECT tally_id FROM evaluation WHERE id=?", (evaluation_id,)
            ).fetchone()
            if not evaluation:
                raise HTTPException(404, "Không tìm thấy lần đánh giá")
            connection.execute(
                f"UPDATE evaluation SET {fields},updated_at=? WHERE id=?",
                (*values.values(), now(), evaluation_id),
            )
            status = payload.status or connection.execute(
                "SELECT status FROM evaluation WHERE id=?", (evaluation_id,)
            ).fetchone()["status"]
            final_status = "completed" if status in {"Chờ phê duyệt", "Hoàn tất"} else "in_progress"
            set_workflow_step(connection, evaluation["tally_id"], "finalize", final_status, {
                "evaluation_status": status,
                "saved_at": now(),
            })
    return {"ok": True}


@router.get("/api/evaluations/{evaluation_id}/report/{kind}")
def report(evaluation_id: str, kind: str):
    conditions = {"important": "important_job=1", "database": "not_in_database=1", "scope": "0"}
    if kind not in conditions:
        raise HTTPException(404, "Báo cáo không tồn tại")
    with db() as connection:
        rows = [
            dict(row)
            for row in connection.execute(
                f"""SELECT * FROM evaluation_item
                WHERE evaluation_id=? AND {conditions[kind]} ORDER BY id""",
                (evaluation_id,),
            )
        ]
    return {
        "items": rows,
        "available": kind != "scope",
        "message": (
            "Chưa có dữ liệu Task–Check/Linked Check để đánh giá tự động."
            if kind == "scope" else None
        ),
    }


def style_sheet(worksheet) -> None:
    blue = "003B71"
    gold = "C89B3C"
    for cell in worksheet[1]:
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column in range(1, worksheet.max_column + 1):
        width = min(
            42,
            max(
                12,
                max(
                    len(str(worksheet.cell(row, column).value or ""))
                    for row in range(1, min(worksheet.max_row, 80) + 1)
                ) + 2,
            ),
        )
        worksheet.column_dimensions[get_column_letter(column)].width = width
    worksheet.sheet_properties.tabColor = gold


EXPORT_COLUMNS = [
    ("Event Number", "task_no"), ("Task Ref (MPD)", "task_ref"), ("Tally Title", "source_title"),
    ("Match Status", "match_status"), ("AMOS Reference", "amos_reference"),
    ("Match Layer", "match_layer"), ("Match Score (%)", "match_score"), ("Match Method", "match_method"),
    ("Candidate Count", "match_candidate_count"),
    ("Suggested Taskcard", "suggested_task_no"),
    ("Suggested Database Title", "suggested_title"), ("Suggested MHR", "suggested_mhr"),
    ("Match Note", "match_note"), ("Taskcard Title", "title"),
    ("Description", "description"), ("Phase", "phase"), ("Zone", "zone"),
    ("Area", "area"), ("Panel", "panel"), ("Special Code", "special_code"),
    ("Resource Type", "resource_type"), ("Classification", "classification"),
    ("Type Code", "type_code"), ("Qualification", "qualification"),
    ("Planning MHR", "planning_mhr"), ("Information Skill", "information_skill"),
    ("Requirement Link", "requirement_link"), ("Linked Jobcards", "linked_jobcards"),
    ("Jobcards Title", "jobcards_title"), ("Important Job", "important_job"),
    ("Not included in Database", "not_in_database"),
]


def export_value(item: dict[str, Any], key: str) -> Any:
    if key == "important_job":
        return "X" if item[key] else None
    if key == "not_in_database":
        if not item[key]:
            return None
        return "X (Đã thêm vào Database)" if item.get("added_to_db") else "X"
    if key == "match_status":
        return STATUS_LABELS.get(item[key], item[key])
    if key == "planning_mhr" and not item["matched"]:
        return None
    return item[key]


@router.get("/api/evaluations/{evaluation_id}/export")
def export_excel(evaluation_id: str):
    with db() as connection:
        header = connection.execute(
            """SELECT e.*,t.* FROM evaluation e JOIN tally_sheet t ON t.id=e.tally_id
            WHERE e.id=?""",
            (evaluation_id,),
        ).fetchone()
        if not header:
            raise HTTPException(404, "Không tìm thấy lần đánh giá")
        items = [
            dict(row)
            for row in connection.execute(
                "SELECT * FROM evaluation_item WHERE evaluation_id=? ORDER BY id",
                (evaluation_id,),
            )
        ]
        progress = workflow_progress(connection, header["tally_id"])
    is_draft = progress["progress"] < 100 or header["status"] not in {"Chờ phê duyệt", "Hoàn tất"}
    workbook = Workbook()
    workbook.remove(workbook.active)
    columns = EXPORT_COLUMNS
    sheets = [
        ("Detailed Task List", items),
        ("Important Jobs", [item for item in items if item["important_job"]]),
        ("Task out of Check Type Scope", []),
        ("Database to be Updated", [item for item in items if item["not_in_database"]]),
    ]
    for title, subset in sheets:
        worksheet = workbook.create_sheet(title)
        worksheet.append([column[0] for column in columns])
        for item in subset:
            worksheet.append([export_value(item, key) for _, key in columns])
        if title == "Task out of Check Type Scope":
            worksheet.append(["Chưa có dữ liệu Task–Check/Linked Check để đánh giá tự động."])
        style_sheet(worksheet)

    worksheet = workbook.create_sheet("FINAL EVALUATION", 0)
    worksheet.append(["BẢN NHÁP - FINAL EVALUATION" if is_draft else "FINAL EVALUATION", ""])
    worksheet.append(["Thông tin", "Giá trị"])
    fields = [
        ("Trạng thái báo cáo", "BẢN NHÁP" if is_draft else header["status"]),
        ("Khách hàng", header["customer"]), ("Mô tả hồ sơ", header["description"]),
        ("Số đăng ký", header["aircraft_reg"]), ("Loại tàu bay", header["aircraft_type"]),
        ("Check Type", header["check_type"]), ("Workpack revision", header["revision"]),
        ("TAT review", header["tat_review"]), ("Tổng task", len(items)),
        ("Task khớp", sum(item["matched"] for item in items)),
        ("Task cần review", sum(item["match_status"] == "REVIEW_REQUIRED" for item in items)),
        ("Task thiếu DB", sum(item["not_in_database"] for item in items)),
        ("Important Job", sum(item["important_job"] for item in items)),
        ("Tổng MHR", round(sum(
            (item["reviewed_mhr"] if item["reviewed_mhr"] is not None else item["planning_mhr"]) or 0
            for item in items
        ), 2)),
        ("Giả định", header["assumptions"]), ("Ghi chú", header["notes"]),
        ("Người lập", header["prepared_by"]), ("Người phê duyệt", header["approved_by"]),
    ]
    for row in fields:
        worksheet.append(row)
    worksheet["A1"].fill = PatternFill(
        "solid", fgColor="B63E3E" if is_draft else "003B71"
    )
    worksheet["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    worksheet.merge_cells("A1:B1")
    worksheet.column_dimensions["A"].width = 28
    worksheet.column_dimensions["B"].width = 60
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    name = f"{'DRAFT_' if is_draft else ''}VAECO_WP_Evaluation_{evaluation_id}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )


@router.get("/api/tally/{tally_id}/download")
def download_tally(tally_id: str):
    with db() as connection:
        row = connection.execute("SELECT * FROM tally_sheet WHERE id=?", (tally_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Không tìm thấy file gốc")
    path = resolve_stored_path(row["stored_path"], row["file_name"])
    if not path.exists():
        raise HTTPException(404, "Không tìm thấy file gốc")
    return FileResponse(path, filename=row["file_name"])


@router.get("/api/evaluations/{evaluation_id}/history")
def history(evaluation_id: str):
    with db() as connection:
        rows = [
            dict(row)
            for row in connection.execute(
                "SELECT * FROM change_log WHERE evaluation_id=? ORDER BY changed_at DESC",
                (evaluation_id,),
            )
        ]
    return {"items": rows}
